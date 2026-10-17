"""Doplnění čísla listu/jízdy do Excelové šablony ZPO (placeholdery {{...}})."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from waste_docs.models import Pruvodka

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # pragma: no cover
    load_workbook = None

ZPO_SND_PLACEHOLDERS = (
    "{{SND_CISLO_LISTU_}}",
    "{{SND_CISLO_LISTU}}",
    "{{CISLO_LISTU_}}",
    "{{CISLO_LISTU}}",
)


def zpo_snd_placeholder_values(pruvodka: Pruvodka) -> dict[str, str]:
    snd = (pruvodka.snd_cislo_listu or "").strip()
    return {token: snd for token in ZPO_SND_PLACEHOLDERS}


def build_zpo_xlsx_bytes(template_file, item: Pruvodka) -> bytes:
    if load_workbook is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")

    with template_file.open("rb") as f:
        raw = f.read()
    workbook = load_workbook(io.BytesIO(raw))
    replacements = zpo_snd_placeholder_values(item)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                if "{{" not in cell.value or "}}" not in cell.value:
                    continue
                output = cell.value
                for token, replacement in replacements.items():
                    output = output.replace(token, replacement)
                cell.value = output

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
