"""List „Poznámky“ v protokolu jizda.xlsx — plný text dlouhých poznámek."""

from __future__ import annotations

from copy import copy

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from waste_docs.models import Pruvodka
from waste_docs.services.reports.protocol_note_overflow import rides_needing_note_supplement

NOTE_SHEET_TITLE = "Poznamky"
NOTE_HEADERS = ("# jízdy", "SND", "RZ", "Řidič", "Subjekt", "Adresa", "Poznámka")


def _driver_label(ride: Pruvodka) -> str:
    snap = (ride.driver_snapshot or "").strip()
    if snap:
        return snap
    if ride.driver_id and ride.driver:
        parts = [ride.driver.first_name or "", ride.driver.last_name or ""]
        return " ".join(p for p in parts if p).strip() or "—"
    return "—"


def _copy_print_setup_from_template(sheet, template_sheet) -> None:
    if template_sheet is None:
        return
    if template_sheet.page_setup:
        sheet.page_setup = copy(template_sheet.page_setup)
    if template_sheet.page_margins:
        sheet.page_margins = copy(template_sheet.page_margins)
    if template_sheet.print_options:
        sheet.print_options = copy(template_sheet.print_options)


def append_protocol_notes_sheet(
    workbook,
    rides: list[Pruvodka],
    *,
    template_sheet=None,
) -> bool:
    """Přidá list s plnými poznámkami; vrátí True pokud byl list vytvořen."""
    supplement_rides = rides_needing_note_supplement(rides)
    if not supplement_rides:
        return False

    if NOTE_SHEET_TITLE in workbook.sheetnames:
        del workbook[NOTE_SHEET_TITLE]

    sheet = workbook.create_sheet(title=NOTE_SHEET_TITLE)
    _copy_print_setup_from_template(sheet, template_sheet)
    header_font = Font(bold=True, size=11)
    wrap = Alignment(wrap_text=True, vertical="top")

    for col_idx, title in enumerate(NOTE_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.alignment = wrap

    for row_idx, ride in enumerate(supplement_rides, start=2):
        values = [
            ride.pk,
            (ride.snd_cislo_listu or "").strip() or "—",
            (ride.vehicle_rz_snapshot or "").strip() or "—",
            _driver_label(ride),
            (ride.name_snapshot or "").strip() or "—",
            (ride.address_snapshot or "").strip() or "—",
            (ride.contact_note or "").strip(),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap
            if col_idx == 7 and value:
                sheet.row_dimensions[row_idx].height = max(30, min(120, 15 * (1 + len(str(value)) // 80)))

    widths = (10, 18, 12, 16, 22, 28, 60)
    for col_idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = width

    return True
