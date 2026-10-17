"""Záložní převod XLSX→PDF bez Excelu/LibreOffice (openpyxl + reportlab)."""

from __future__ import annotations

from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


class BasicXlsxPdfError(RuntimeError):
    pass


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _trim_rows(rows: list[list[str]]) -> list[list[str]]:
    while rows and all(not cell for cell in rows[-1]):
        rows.pop()
    if not rows:
        return [[""]]
    max_cols = max(len(row) for row in rows)
    trimmed_cols = max_cols
    while trimmed_cols > 0:
        col_idx = trimmed_cols - 1
        if any(col_idx < len(row) and row[col_idx] for row in rows):
            break
        trimmed_cols -= 1
    if trimmed_cols <= 0:
        return [[""]]
    out: list[list[str]] = []
    for row in rows:
        padded = row[:trimmed_cols] + [""] * max(0, trimmed_cols - len(row))
        out.append(padded[:trimmed_cols])
    return out


def _iter_sheet_rows(sheet) -> list[list[str]]:
    rows: list[list[str]] = []
    empty_streak = 0
    for row in sheet.iter_rows(values_only=True):
        cells = [_cell_text(value) for value in row]
        if not any(cells):
            empty_streak += 1
            if empty_streak >= 8 and rows:
                break
            continue
        empty_streak = 0
        rows.append(cells)
    return _trim_rows(rows)


def convert_xlsx_to_basic_pdf(source: Path, dest: Path) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise BasicXlsxPdfError("Chybí openpyxl pro záložní převod XLSX→PDF.") from exc

    if not source.is_file():
        raise BasicXlsxPdfError("Soubor neexistuje.")

    dest.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows = _iter_sheet_rows(sheet)
    finally:
        workbook.close()

    if len(rows) > 250 or (rows and len(rows[0]) > 40):
        raise BasicXlsxPdfError(
            "Soubor je příliš velký pro záložní PDF převod — nainstalujte LibreOffice."
        )

    page_size = landscape(A4) if rows and len(rows[0]) > 8 else A4
    doc = SimpleDocTemplate(
        str(dest),
        pagesize=page_size,
        leftMargin=8 * mm,
        rightMargin=8 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
    )
    body_style = ParagraphStyle(
        name="Cell",
        fontName="Helvetica",
        fontSize=7,
        leading=8,
    )
    table_data = [
        [Paragraph(cell.replace("\n", "<br/>") or " ", body_style) for cell in row]
        for row in rows
    ]
    col_count = len(rows[0])
    col_width = (doc.width - 4) / max(col_count, 1)
    table = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    doc.build(
        [
            Paragraph(
                f"Záložní tisk — {source.name}",
                ParagraphStyle(name="Title", fontName="Helvetica-Bold", fontSize=8),
            ),
            Spacer(1, 4),
            table,
        ]
    )
