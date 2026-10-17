"""Klon listu protokolu — vždy z čisté šablony na disku (ne copy_worksheet v sešitu)."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import uuid4

from waste_docs.services.reports.protocol_sheet_bounds import trim_protocol_sheet_to_bounds

try:
    from openpyxl import load_workbook
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError:  # pragma: no cover
    load_workbook = None  # type: ignore[assignment,misc]
    Worksheet = None  # type: ignore[assignment,misc]


class ProtocolTemplateSource:
    """Drží bytes šablony a referenční list pro tokeny a fill."""

    def __init__(self, template_path: Path, *, keep_vba: bool) -> None:
        if load_workbook is None:
            raise RuntimeError("Knihovna openpyxl není nainstalovaná.")
        self.template_path = Path(template_path)
        self.keep_vba = keep_vba
        self.template_bytes = self.template_path.read_bytes()
        self._ref_wb = load_workbook(BytesIO(self.template_bytes), keep_vba=keep_vba)
        self.template_sheet = self._ref_wb.active
        trim_protocol_sheet_to_bounds(self.template_sheet)

    def close(self) -> None:
        self._ref_wb.close()

    def clone_sheet(
        self,
        workbook,
        *,
        copy_images_fn=None,
    ) -> Worksheet:
        """Přidá nový list zkopírovaný z nezměněné šablony."""
        donor_wb = load_workbook(BytesIO(self.template_bytes), keep_vba=self.keep_vba)
        donor_ws = donor_wb.active
        trim_protocol_sheet_to_bounds(donor_ws)
        donor_title = f"__donor_{uuid4().hex[:8]}"
        donor_ws.title = donor_title

        donor_ws._parent = workbook
        workbook._sheets.append(donor_ws)
        try:
            new_sheet = workbook.copy_worksheet(donor_ws)
            if copy_images_fn is not None:
                copy_images_fn(donor_ws, new_sheet)
            trim_protocol_sheet_to_bounds(new_sheet)
            del workbook[donor_title]
            return new_sheet
        finally:
            if donor_ws in workbook._sheets:
                workbook._sheets.remove(donor_ws)
            donor_ws._parent = donor_wb
            donor_wb.close()


def open_protocol_output_workbook(template_source: ProtocolTemplateSource):
    """Výstupní sešit se zachovaným VBA ze šablony (makro TiskCelySesit)."""
    if load_workbook is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")
    workbook = load_workbook(
        BytesIO(template_source.template_bytes),
        keep_vba=template_source.keep_vba,
    )
    trim_protocol_sheet_to_bounds(workbook.active)
    return workbook
