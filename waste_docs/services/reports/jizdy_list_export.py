"""Export seznamu jízd do XLSX — sloupce definované zde (jedno místo pro údržbu)."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from io import BytesIO
from typing import Any, Callable

from django.utils import timezone

from waste_docs.follow_up_display import follow_up_child_link_symbol, follow_up_parent_link_symbol
from waste_docs.services.pruvodka_list_filters import PruvodkaListFilterState

try:
    from openpyxl import Workbook
except ModuleNotFoundError:  # pragma: no cover
    Workbook = None

MAX_EXPORT_ROWS = 10_000


@dataclass(frozen=True)
class _Col:
    header: str
    get_value: Callable[[Any, "JizdyExportContext"], str]


@dataclass
class JizdyExportContext:
    awaiting_follow_up: bool
    today: date


def _user_label(u) -> str:
    if not u:
        return "-"
    return (u.get_full_name() or u.username or "").strip() or "-"


def _follow_up_text(item, _ctx: JizdyExportContext) -> str:
    parts: list[str] = []
    if item.follow_up_open:
        parts.append("čeká")
    if item.follow_up_parent_id:
        parts.append(f"{follow_up_parent_link_symbol(item)}{item.follow_up_parent_id}")
    children = list(item.follow_up_children.all())
    if children:
        parts.append(follow_up_child_link_symbol(item) + ",".join(str(c.pk) for c in children))
    return " ".join(parts) if parts else "—"


def _creator_dv_text(item, _ctx: JizdyExportContext) -> str:
    main = _user_label(item.created_by)
    dv_labels = [_user_label(st.created_by) for st in item.svoz_task_sources.all()]
    if not dv_labels:
        return main
    return f"{main} | Žádost DV: " + "; ".join(dv_labels)


def _containers(item, _ctx: JizdyExportContext) -> str:
    a = (item.container_number or "").strip() or "-"
    b = (item.container_number_2 or "").strip()
    if b:
        return f"{a} / {b}"
    return a


def _days_on_site(item, ctx: JizdyExportContext) -> str:
    n = max(0, (ctx.today - item.drive_date).days)
    return str(n)


def _columns(awaiting: bool) -> list[_Col]:
    cols: list[_Col] = [
        _Col("ID", lambda i, c: str(i.pk)),
        _Col("Číslo listu/jízdy", lambda i, c: (i.snd_cislo_listu or "").strip() or "-"),
        _Col("Status", lambda i, c: i.get_status_display()),
        _Col("Návaznost", _follow_up_text),
        _Col("Datum jízdy", lambda i, c: i.drive_date.strftime("%d.%m.%Y")),
    ]
    if awaiting:
        cols.append(_Col("Na místě (dny)", _days_on_site))
    cols.extend(
        [
            _Col("RZ", lambda i, c: (i.vehicle_rz_snapshot or "").strip() or "-"),
            _Col("Řidič", lambda i, c: (i.driver_snapshot or "").strip() or "-"),
            _Col("Subjekt", lambda i, c: (i.name_snapshot or "").strip() or "-"),
            _Col("Kontejner (1) / (2)", _containers),
            _Col("Koncovka", lambda i, c: (i.suffix_snapshot or "").strip() or "-"),
            _Col("Číslo smlouvy", lambda i, c: (i.contract_number_snapshot or "").strip() or "-"),
            _Col("Telefon", lambda i, c: (i.contact_phone or "").strip() or "-"),
            _Col("Vytvořil / žádost DV", _creator_dv_text),
        ]
    )
    return cols


def build_jizdy_list_xlsx_bytes(queryset, state: PruvodkaListFilterState) -> tuple[bytes, str]:
    if Workbook is None:  # pragma: no cover
        raise RuntimeError("openpyxl není nainstalován.")

    today = timezone.localdate()
    ctx = JizdyExportContext(awaiting_follow_up=state.awaiting_follow_up, today=today)
    cols = _columns(state.awaiting_follow_up)

    rows_list = list(queryset[: MAX_EXPORT_ROWS + 1])
    truncated = len(rows_list) > MAX_EXPORT_ROWS
    if truncated:
        rows_list = rows_list[:MAX_EXPORT_ROWS]

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Jízdy"

    row_num = 1
    if truncated:
        ws.cell(row=row_num, column=1, value=f"Export omezen na {MAX_EXPORT_ROWS} řádků (zbytek neexportován).")
        row_num += 1

    for c_idx, col in enumerate(cols, start=1):
        ws.cell(row=row_num, column=c_idx, value=col.header)
    row_num += 1

    for item in rows_list:
        for c_idx, col in enumerate(cols, start=1):
            ws.cell(row=row_num, column=c_idx, value=col.get_value(item, ctx))
        row_num += 1

    buf = BytesIO()
    wb.save(buf)

    fname = "jizdy-export"
    if state.selected_date:
        fname += f"_{state.selected_date.isoformat()}"
    else:
        fname += "_vsechny"
    if state.awaiting_follow_up:
        fname += "_navaznost"
    fname += ".xlsx"
    return buf.getvalue(), fname
