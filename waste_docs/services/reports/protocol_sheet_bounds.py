"""Hranice listu protokolu — layout A1:M38 (jedna stránka A4 na šířku)."""

from __future__ import annotations

from typing import Iterator

from waste_docs.services.reports.protocol_page_setup import PROTOCOL_SHEET_PRINT_AREA

PROTOCOL_SHEET_MAX_ROW = 38
PROTOCOL_SHEET_MAX_COL = 13  # M


def iter_protocol_cells(sheet) -> Iterator:
    """Buňky protokolu — jen řádky 1–38, sloupce A–M."""
    for row in sheet.iter_rows(
        min_row=1,
        max_row=PROTOCOL_SHEET_MAX_ROW,
        min_col=1,
        max_col=PROTOCOL_SHEET_MAX_COL,
    ):
        for cell in row:
            yield cell


def trim_protocol_sheet_to_bounds(sheet) -> None:
    """Odstraní ghost řádky/sloupce pod layoutem protokolu (zbytek Excel tabulky apod.)."""
    max_row = sheet.max_row or 0
    if max_row > PROTOCOL_SHEET_MAX_ROW:
        sheet.delete_rows(PROTOCOL_SHEET_MAX_ROW + 1, max_row - PROTOCOL_SHEET_MAX_ROW)

    max_col = sheet.max_column or 0
    if max_col > PROTOCOL_SHEET_MAX_COL:
        sheet.delete_cols(PROTOCOL_SHEET_MAX_COL + 1, max_col - PROTOCOL_SHEET_MAX_COL)

    for row_num in list(sheet.row_dimensions.keys()):
        if row_num > PROTOCOL_SHEET_MAX_ROW:
            del sheet.row_dimensions[row_num]

    for col_letter in list(sheet.column_dimensions.keys()):
        from openpyxl.utils import column_index_from_string

        if column_index_from_string(col_letter) > PROTOCOL_SHEET_MAX_COL:
            del sheet.column_dimensions[col_letter]

    sheet.calculate_dimension()
    sheet.print_area = PROTOCOL_SHEET_PRINT_AREA
