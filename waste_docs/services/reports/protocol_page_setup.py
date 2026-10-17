"""Nastavení stránky protokolu — jednostranný tisk A1:M38 (1 jízda = 1 list = 1 stránka).

Starší duplex oblasti A1:M39 + A40:M79 ponechány pro zpětnou kompatibilitu oprav existujících souborů.
"""

from __future__ import annotations

import shutil
import tempfile
import uuid
from copy import copy
from pathlib import Path

from pypdf import PdfReader, PdfWriter

from waste_docs.services.reports.protocol_note_sheet import NOTE_SHEET_TITLE

PROTOCOL_SHEET_PRINT_AREA = "A1:M38"
PROTOCOL_SHEET_PRINT_AREA_FRONT = "A1:M39"
PROTOCOL_SHEET_PRINT_AREA_BACK = "A40:M79"
PROTOCOL_SHEET_PRINT_AREA_FULL = "A1:M79"
PROTOCOL_SHEET_ROW_BREAK = 40
PROTOCOL_SHEET_PRINT_AREAS = (
    PROTOCOL_SHEET_PRINT_AREA_FRONT,
    PROTOCOL_SHEET_PRINT_AREA_BACK,
)
PROTOCOL_SHEET_EXCEL_PRINT_AREA = (
    f"{PROTOCOL_SHEET_PRINT_AREA_FRONT},{PROTOCOL_SHEET_PRINT_AREA_BACK}"
)

try:
    from openpyxl import load_workbook
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.properties import PageSetupProperties
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError:
    load_workbook = None  # type: ignore[assignment,misc]
    Break = None  # type: ignore[assignment,misc]
    PageSetupProperties = None  # type: ignore[assignment,misc]
    Worksheet = None  # type: ignore[assignment,misc]


def _apply_pdf_block_page_setup(ws: Worksheet, print_area: str) -> None:
    """Jedna print oblast → 1 PDF stránka (fit 1×1, landscape)."""
    if PageSetupProperties is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")
    if ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties()
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.orientation = "landscape"
    ws.print_area = print_area


def apply_single_side_protocol_page_setup(ws: Worksheet, template_sheet: Worksheet) -> None:
    """Jednostranný tisk: page setup ze šablony, print area jen A1:M38."""
    if PageSetupProperties is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")
    if template_sheet.page_setup:
        ws.page_setup = copy(template_sheet.page_setup)
    if template_sheet.page_margins:
        ws.page_margins = copy(template_sheet.page_margins)
    if template_sheet.print_options:
        ws.print_options = copy(template_sheet.print_options)
    if template_sheet.sheet_properties.pageSetUpPr is not None:
        ws.sheet_properties.pageSetUpPr = copy(template_sheet.sheet_properties.pageSetUpPr)
    elif ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties()
    ws.print_area = PROTOCOL_SHEET_PRINT_AREA
    ws.row_breaks.brk = []


def apply_protocol_page_setup_from_template(ws: Worksheet, template_sheet: Worksheet) -> None:
    """Excel tisk: page setup ze šablony + duplex print oblasti (bez row break 40)."""
    if PageSetupProperties is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")
    if template_sheet.page_setup:
        ws.page_setup = copy(template_sheet.page_setup)
    if template_sheet.page_margins:
        ws.page_margins = copy(template_sheet.page_margins)
    if template_sheet.print_options:
        ws.print_options = copy(template_sheet.print_options)
    if template_sheet.sheet_properties.pageSetUpPr is not None:
        ws.sheet_properties.pageSetUpPr = copy(template_sheet.sheet_properties.pageSetUpPr)
    elif ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties()
    ws.print_area = PROTOCOL_SHEET_EXCEL_PRINT_AREA
    ws.row_breaks.brk = []


def apply_protocol_page_setup_to_worksheet(ws: Worksheet, template_sheet: Worksheet | None = None) -> None:
    """Zpětná kompatibilita — preferuje page setup ze šablony."""
    if template_sheet is not None:
        apply_protocol_page_setup_from_template(ws, template_sheet)
        return
    if PageSetupProperties is None or Break is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")
    if ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties()
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 2
    ws.page_setup.orientation = "landscape"
    ws.print_area = PROTOCOL_SHEET_EXCEL_PRINT_AREA
    ws.row_breaks.brk = []
    ws.row_breaks.append(Break(id=PROTOCOL_SHEET_ROW_BREAK))


def apply_protocol_page_setup_to_workbook(workbook, template_sheet: Worksheet | None = None) -> None:
    for ws in workbook.worksheets:
        if ws.title == NOTE_SHEET_TITLE:
            continue
        apply_protocol_page_setup_to_worksheet(ws, template_sheet=template_sheet)


def ensure_protocol_page_setup(xlsx_path: Path, template_path: Path | None = None) -> None:
    """Uloží page setup do existujícího protokolu (podle šablony jizda.xlsx)."""
    if load_workbook is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")
    path = Path(xlsx_path)
    if not path.is_file():
        return
    from waste_docs.services.reports.protocol_print_settings import resolve_protocol_template_path

    template_file = resolve_protocol_template_path(template_path)
    template_wb = load_workbook(template_file)
    try:
        template_sheet = template_wb.active
        workbook = load_workbook(path)
        try:
            apply_protocol_page_setup_to_workbook(workbook, template_sheet=template_sheet)
            workbook.save(path)
        finally:
            workbook.close()
    finally:
        template_wb.close()


def ensure_protocol_template_page_setup(template_path: Path) -> None:
    """Opraví page setup ve šabloně jizda.xlsx (listy protokolu)."""
    ensure_protocol_page_setup(template_path)


def _convert_single_sheet_area_to_pdf(
    source_path: Path,
    sheet_title: str,
    print_area: str,
    temp_dir: Path,
) -> Path:
    from waste_docs.services.reports.xlsx_to_pdf import XlsxToPdfError, _convert_with_libreoffice

    if load_workbook is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")

    from waste_docs.services.reports.protocol_pdf_fonts import normalize_worksheet_fonts

    workbook = load_workbook(source_path)
    try:
        for name in list(workbook.sheetnames):
            if name != sheet_title:
                del workbook[name]
        worksheet = workbook[sheet_title]
        _apply_pdf_block_page_setup(worksheet, print_area)
        normalize_worksheet_fonts(worksheet)
        temp_xlsx = temp_dir / f"area_{uuid.uuid4().hex}.xlsx"
        temp_pdf = temp_dir / f"area_{uuid.uuid4().hex}.pdf"
        workbook.save(temp_xlsx)
    finally:
        workbook.close()

    try:
        _convert_with_libreoffice(temp_xlsx, temp_pdf)
    except Exception as exc:
        raise XlsxToPdfError(
            f"Převod bloku protokolu ({print_area}) selhal: {exc}"
        ) from exc
    finally:
        temp_xlsx.unlink(missing_ok=True)
    return temp_pdf


def convert_protocol_xlsx_to_pdf(source: Path, dest: Path) -> None:
    """Každý list protokolu = 1 PDF stránka (A1:M38, landscape, fit 1×1)."""
    if load_workbook is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")

    source = Path(source)
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(source, read_only=True)
    try:
        sheet_titles = [
            name
            for name in workbook.sheetnames
            if name != NOTE_SHEET_TITLE
        ]
    finally:
        workbook.close()

    if not sheet_titles:
        raise RuntimeError("Protokol neobsahuje žádný list k tisku.")

    writer = PdfWriter()
    temp_dir = Path(tempfile.mkdtemp(prefix="jidis_protocol_pdf_"))
    try:
        for sheet_title in sheet_titles:
            temp_pdf = _convert_single_sheet_area_to_pdf(
                source,
                sheet_title,
                PROTOCOL_SHEET_PRINT_AREA,
                temp_dir,
            )
            reader = PdfReader(str(temp_pdf))
            for page in reader.pages:
                writer.add_page(page)
            temp_pdf.unlink(missing_ok=True)
        with dest.open("wb") as output:
            writer.write(output)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def is_protocol_xlsx_path(source: Path | str) -> bool:
    """True pokud jde o vygenerovaný protokol v media/protokoly/."""
    from django.conf import settings

    try:
        media_root = Path(settings.MEDIA_ROOT).resolve()
        rel = Path(source).resolve().relative_to(media_root).as_posix()
    except (ValueError, OSError, TypeError):
        return False
    return rel.replace("\\", "/").startswith("protokoly/") and rel.lower().endswith(".xlsx")
