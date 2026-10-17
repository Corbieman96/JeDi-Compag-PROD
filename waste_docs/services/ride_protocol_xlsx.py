from __future__ import annotations

import re
import unicodedata
from copy import copy, deepcopy
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from uuid import uuid4

from django.conf import settings
from django.contrib.auth.models import User

from catalogs.models import VehicleRzProfile
from users.services.access import can_see_all_waste_docs
from waste_docs.services.ride_print_order import order_rides_queryset
from waste_docs.models import Pruvodka
from waste_docs.activity_fields import activity_uses_fulfillment_address, protocol_address_snapshot_for_ride
from waste_docs.services.reports.protocol_note_overflow import protocol_note_for_cell
from waste_docs.services.reports.protocol_note_sheet import append_protocol_notes_sheet
from waste_docs.services.reports.protocol_fill_sheet import (
    PROTOCOL_KDE_LABEL_COORDS,
    fill_protocol_sheet,
)
from waste_docs.services.reports.protocol_poznamka_highlight import apply_poznamka_highlight
from waste_docs.services.reports.protocol_print_settings import (
    ensure_protocol_simplex_print_settings,
    resolve_protocol_template_path,
    should_keep_vba_for_template,
    validate_generated_xlsx,
)

from waste_docs.services.reports.protocol_sheet_bounds import iter_protocol_cells
from waste_docs.services.reports.protocol_template_sheet import (
    ProtocolTemplateSource,
    open_protocol_output_workbook,
)

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # pragma: no cover
    load_workbook = None

PROTOCOL_LINES_PER_SHEET = 7
# Statické popisky ve jizda.xlsx — regresní testy ověřují beze změny.
PROTOCOL_STATIC_LABEL_COORDS = ("C1", "G2", "B4", "F4", "H4")


TOKEN_RE = re.compile(r"\{\{[^{}]+\}\}")
INDEXED_TOKEN_RE = re.compile(r"^\{\{([A-Z0-9_]+)_(\d+)\}\}$")
SND_BASE_RE = re.compile(r"^(.*)-\d+$")

LINE_FIELD_MAP = {
    "SND_CISLO_LISTU": "snd_cislo_listu",
    "CINNOST": "activity_snapshot",
    "KDO": "name_snapshot",
    "MISTO_DODANI_CINNOST": "address_snapshot",
    "DRUH_ODPADU": "waste_type_snapshot",
    "KONCOVKA": "suffix_snapshot",
    "KOD": "waste_code_snapshot",
    "KATEGORIE": "waste_category_snapshot",
    "KONTEJNER": "container_number",
    "OBJEDNAVKA": "contract_number_snapshot",
    "POZNAMKA": "contact_note",
    "TELEFON": "contact_phone",
    "KONTAKT": "contact_phone",
    "NADOBA": "container_number_2",
    "VELIKOST": "container_number_2",
}


def _driver_last_name(ride: Pruvodka) -> str:
    if ride.driver_id and ride.driver and (ride.driver.last_name or "").strip():
        return (ride.driver.last_name or "").strip()
    snap = (ride.driver_snapshot or "").strip()
    return snap.split()[-1] if snap else "Bez řidiče"


def _vehicle_kind_label(vehicle_rz_id: int | None) -> str:
    if not vehicle_rz_id:
        return ""
    profile = (
        VehicleRzProfile.objects.filter(rz_id=vehicle_rz_id)
        .prefetch_related("vehicle_kinds")
        .first()
    )
    if not profile:
        return ""
    kinds = list(profile.vehicle_kinds.all())
    if kinds:
        return ", ".join(k.name for k in kinds)
    return (profile.vehicle_type or "").strip()


def _snd_base(value: str) -> str:
    text = (value or "").strip()
    if not text:
        return ""
    m = SND_BASE_RE.match(text)
    return m.group(1) if m else text


def _snd_group_short(value: str) -> str:
    """Číslo SND skupiny bez úvodních nul (0000493-2026-1 → 493) pro hlavičku protokolu."""
    text = (value or "").strip()
    if not text:
        return ""
    segment = text.split("-", 1)[0]
    try:
        return str(int(segment))
    except ValueError:
        stripped = segment.lstrip("0")
        return stripped if stripped else segment or "0"


def _page_snd_header(lines: list[dict], *, max_len: int = 72) -> str:
    """Hlavička SND: jedna skupina = plné číslo (0000295-2026), mix = zkráceně (493, 471)."""
    seen_bases: set[str] = set()
    bases_ordered: list[str] = []
    for line in lines:
        base = _snd_base(str(line.get("snd_cislo_listu") or ""))
        if not base or base in seen_bases:
            continue
        seen_bases.add(base)
        bases_ordered.append(base)
    if not bases_ordered:
        return ""
    if len(bases_ordered) == 1:
        return bases_ordered[0]

    groups = [_snd_group_short(base) for base in bases_ordered]
    full = ", ".join(groups)
    if len(full) <= max_len:
        return full

    included: list[str] = []
    for g in groups:
        trial_list = included + [g]
        remaining = len(groups) - len(trial_list)
        trial = ", ".join(trial_list)
        trial_with_suffix = f"{trial} +{remaining}" if remaining else trial
        if len(trial_with_suffix) <= max_len:
            included.append(g)
        else:
            break

    if not included:
        remaining = len(groups) - 1
        suffix = f" +{remaining}"
        first = groups[0]
        if len(first) + len(suffix) <= max_len:
            return f"{first}{suffix}"
        return f"{first[: max(1, max_len - len(suffix))]}{suffix}"

    remaining = len(groups) - len(included)
    base = ", ".join(included)
    if remaining:
        return f"{base} +{remaining}"
    return base


def _safe_sheet_title(base_title: str, used: set[str]) -> str:
    ascii_title = (
        unicodedata.normalize("NFKD", (base_title or "").strip())
        .encode("ascii", "ignore")
        .decode("ascii")
    )
    ascii_title = re.sub(r"\s+", " ", ascii_title).strip()
    cleaned = re.sub(r"[:\\/?*\[\]]", "-", ascii_title) or "Protokol"
    cleaned = cleaned[:31]
    if cleaned not in used:
        used.add(cleaned)
        return cleaned
    idx = 2
    while True:
        suffix = f" ({idx})"
        candidate = f"{cleaned[: max(1, 31 - len(suffix))]}{suffix}"
        if candidate not in used:
            used.add(candidate)
            return candidate
        idx += 1


def _collect_template_tokens(sheet) -> set[str]:
    tokens: set[str] = set()
    for cell in iter_protocol_cells(sheet):
        if isinstance(cell.value, str) and "{{" in cell.value and "}}" in cell.value:
            tokens.update(TOKEN_RE.findall(cell.value))
    return tokens


def _line_value(token_base: str, line: dict) -> str:
    attr = LINE_FIELD_MAP.get(token_base)
    if not attr:
        return ""
    return str((line.get(attr) or "")).strip()


def _token_value(token: str, page: dict) -> str:
    if token == "{{DRUH}}":
        return page["vehicle_kind"]
    if token == "{{DATUM}}":
        return page["drive_date"]
    if token == "{{SND}}":
        return page["snd_base"]
    if token == "{{RZ_VOZIDLA}}":
        return page["vehicle_rz"]
    if token == "{{RIDIC}}":
        return page["driver_last_name"]
    if token == "{{PRIVES}}":
        return ""

    m = INDEXED_TOKEN_RE.match(token)
    if not m:
        return ""
    token_base, idx_raw = m.groups()
    idx = int(idx_raw)
    if idx < 1 or idx > 7:
        return ""
    lines = page["lines"]
    line = lines[idx - 1] if idx - 1 < len(lines) else {}
    return _line_value(token_base, line)


def _copy_cell_font_from_template(template_sheet, output_sheet, coord: str) -> None:
    tpl_cell = template_sheet[coord]
    out_cell = output_sheet[coord]
    if tpl_cell.font:
        out_cell.font = copy(tpl_cell.font)


def _render_sheet_tokens(sheet, page: dict, tokens: set[str], template_sheet=None) -> None:
    for cell in iter_protocol_cells(sheet):
        if not isinstance(cell.value, str) or "{{" not in cell.value or "}}" not in cell.value:
            continue
        original = cell.value
        text = original
        for token in tokens:
            if token in text:
                text = text.replace(token, _token_value(token, page))
        if text == original:
            continue
        cell.value = text
        if template_sheet is not None:
            _copy_cell_font_from_template(template_sheet, sheet, cell.coordinate)


def _copy_template_images(template_sheet, target_sheet) -> None:
    for image in getattr(template_sheet, "_images", []) or []:
        try:
            cloned = deepcopy(image)
            cloned.anchor = deepcopy(image.anchor)
            target_sheet.add_image(cloned)
        except Exception:
            # Pokud knihovna nepodporuje klon daného objektu obrázku, pokračujeme bez pádu.
            continue


def _chunked(items: list, size: int) -> list[list]:
    return [items[i : i + size] for i in range(0, len(items), size)]


def _safe_filename_token(raw: str, fallback: str) -> str:
    token = re.sub(r"[^0-9A-Za-z._-]+", "-", (raw or "").strip()).strip("-._")
    return token or fallback


PROTOCOL_SCOPE_DAY = "day"
PROTOCOL_SCOPE_GROUP = "group"
PROTOCOL_SCOPE_SINGLE = "single"


def _legacy_protocol_filename_token(rides: list[Pruvodka]) -> str:
    """Starší pojmenování podle SND první jízdy — kvůli dohledání existujících souborů."""
    if not rides:
        return "protokol"
    if len(rides) == 1:
        return _safe_filename_token(rides[0].snd_cislo_listu or "", "protokol")
    first = rides[0]
    snd_base = _snd_base(first.snd_cislo_listu or "")
    if snd_base:
        return _safe_filename_token(snd_base, "protokol")
    if first.drive_date:
        return _safe_filename_token(first.drive_date.isoformat(), "protokol")
    return "protokol"


def protocol_filename_token(
    rides: list[Pruvodka],
    *,
    scope: str = PROTOCOL_SCOPE_SINGLE,
    selected_date: date | None = None,
) -> str:
    if scope == PROTOCOL_SCOPE_DAY:
        drive_date = selected_date or (rides[0].drive_date if rides else None)
        if drive_date:
            return _safe_filename_token(f"den_{drive_date.isoformat()}", "protokol")
    if not rides:
        return "protokol"
    if scope == PROTOCOL_SCOPE_GROUP:
        drive_date = rides[0].drive_date
        rz_id = rides[0].vehicle_rz_id or 0
        driver_id = rides[0].driver_id or 0
        if drive_date:
            return _safe_filename_token(
                f"skupina_{drive_date.isoformat()}_rz{rz_id}_rid{driver_id}",
                "protokol",
            )
    return _legacy_protocol_filename_token(rides)


def _protocol_filename_token(rides: list[Pruvodka]) -> str:
    return protocol_filename_token(rides, scope=PROTOCOL_SCOPE_SINGLE)


def protocol_filename_tokens_for_lookup(
    rides: list[Pruvodka],
    *,
    scope: str,
    selected_date: date | None = None,
) -> list[str]:
    tokens = [protocol_filename_token(rides, scope=scope, selected_date=selected_date)]
    if scope in (PROTOCOL_SCOPE_DAY, PROTOCOL_SCOPE_GROUP):
        legacy = _legacy_protocol_filename_token(rides)
        if legacy not in tokens:
            tokens.append(legacy)
    if scope == PROTOCOL_SCOPE_GROUP and rides:
        legacy_rz = _safe_filename_token(
            f"skupina_{rides[0].drive_date.isoformat()}_rz{rides[0].vehicle_rz_id or 0}",
            "protokol",
        )
        if legacy_rz not in tokens:
            tokens.append(legacy_rz)
    return tokens


def protocol_latest_ride_timestamp(rides: list[Pruvodka]) -> float:
    latest = 0.0
    for ride in rides:
        for ts in (getattr(ride, "created_at", None), getattr(ride, "updated_at", None)):
            if not ts:
                continue
            if hasattr(ts, "timestamp"):
                latest = max(latest, ts.timestamp())
            else:
                latest = max(latest, datetime.combine(ts, datetime.min.time()).timestamp())
    return latest


def protocol_file_is_stale(
    rides: list[Pruvodka],
    path: Path,
    *,
    plan_date: date | None = None,
) -> bool:
    if not path.exists():
        return True
    try:
        mtime = path.stat().st_mtime
    except OSError:
        return True
    if rides and mtime < protocol_latest_ride_timestamp(rides):
        return True
    if plan_date is not None:
        from waste_docs.services.reports.protocol_plan_slots import (
            plan_only_protocol_slots,
            plan_tasks_changed_since,
        )

        if plan_only_protocol_slots(plan_date) and plan_tasks_changed_since(plan_date, mtime):
            return True
        if not rides and plan_only_protocol_slots(plan_date):
            return False
    return False


def _vehicle_group_key(ride: Pruvodka) -> tuple[int, str]:
    """Odděluje jízdy podle vozidla (RZ), aby se do jednoho listu nedostaly dva vozy u stejného řidiče/SND."""
    return (ride.vehicle_rz_id or 0, (ride.vehicle_rz_snapshot or "").strip())


def _line_dict(ride: Pruvodka) -> dict:
    note_raw = (ride.contact_note or "").strip()
    note_cell, _needs_supplement = protocol_note_for_cell(note_raw)
    uses_address = activity_uses_fulfillment_address(ride.activity_id)
    return {
        "snd_cislo_listu": (ride.snd_cislo_listu or "").strip(),
        "activity_snapshot": (ride.activity_snapshot or "").strip(),
        "name_snapshot": (ride.name_snapshot or "").strip(),
        "address_snapshot": protocol_address_snapshot_for_ride(ride.address_snapshot, ride.activity_id),
        "suppress_kde": not uses_address,
        "waste_type_snapshot": (ride.waste_type_snapshot or "").strip(),
        "suffix_snapshot": (ride.suffix_snapshot or "").strip(),
        "waste_code_snapshot": (ride.waste_code_snapshot or "").strip(),
        "waste_category_snapshot": (ride.waste_category_snapshot or "").strip(),
        "container_number": (ride.container_number or "").strip(),
        "contract_number_snapshot": (ride.contract_number_snapshot or "").strip(),
        "contact_note": note_cell,
        "contact_phone": (ride.contact_phone or "").strip(),
        "container_number_2": (ride.container_number_2 or "").strip(),
    }


def rides_for_protocol_group(
    selected_date: date,
    vehicle_rz_id: int,
    user: User,
    *,
    filter_pks: list[int] | None = None,
    for_email_package: bool = False,
    driver_id: int | None = None,
) -> list[Pruvodka]:
    from waste_docs.services.ride_document_email.exclusions import is_group_emailed

    if not for_email_package and is_group_emailed(selected_date, vehicle_rz_id):
        return []
    qs = order_rides_queryset(
        Pruvodka.objects.filter(drive_date=selected_date, vehicle_rz_id=vehicle_rz_id, is_archived=False).select_related(
            "driver", "vehicle_rz"
        )
    )
    if driver_id is not None and not filter_pks:
        if driver_id > 0:
            qs = qs.filter(driver_id=driver_id)
        else:
            qs = qs.filter(driver_id__isnull=True)
    if filter_pks:
        qs = qs.filter(pk__in=filter_pks)
    if not can_see_all_waste_docs(user):
        qs = qs.filter(created_by=user)
    return list(qs)


def rides_for_protocol_day(selected_date: date, user: User) -> list[Pruvodka]:
    from waste_docs.services.ride_document_email.exclusions import emailed_vehicle_rz_ids_for_date

    emailed = emailed_vehicle_rz_ids_for_date(selected_date)
    qs = order_rides_queryset(
        Pruvodka.objects.filter(drive_date=selected_date, is_archived=False).select_related("driver", "vehicle_rz")
    )
    if emailed:
        qs = qs.exclude(vehicle_rz_id__in=emailed)
    if not can_see_all_waste_docs(user):
        qs = qs.filter(created_by=user)
    return list(qs)


def rides_for_protocol_single(ride_pk: int, user: User) -> list[Pruvodka]:
    qs = Pruvodka.objects.filter(pk=ride_pk, is_archived=False).select_related("driver", "vehicle_rz")
    if not can_see_all_waste_docs(user):
        qs = qs.filter(created_by=user)
    return list(qs)


def _group_page_header(driver_rides: list[Pruvodka], driver_name: str) -> dict:
    vehicle_kind = _vehicle_kind_label(driver_rides[0].vehicle_rz_id)
    drive_date = (
        driver_rides[0].drive_date.strftime("%d.%m.%Y") if driver_rides[0].drive_date else ""
    )
    rz = (driver_rides[0].vehicle_rz_snapshot or "").strip()
    sheet_bits = [driver_name]
    if rz:
        sheet_bits.append(rz)
    return {
        "driver_last_name": driver_name,
        "vehicle_kind": vehicle_kind,
        "drive_date": drive_date,
        "vehicle_rz": rz,
        "sheet_title_base": " ".join(sheet_bits) or "Protokol",
    }


def build_protocol_pages(
    rides: list[Pruvodka],
    plan_slots=None,
    *,
    vehicle_kind_label_fn=_vehicle_kind_label,
) -> list[dict]:
    """Dynamicky sestaví seznam listů: až 7 jízd/list, seskupeno podle řidič+RZ."""
    plan_slots = plan_slots or []
    rides_by_group: dict[tuple[str, int, str], list[Pruvodka]] = defaultdict(list)
    for ride in rides:
        driver_name = _driver_last_name(ride)
        v_key = _vehicle_group_key(ride)
        rides_by_group[(driver_name, v_key[0], v_key[1])].append(ride)

    sort_key = lambda k: (k[0].lower(), k[1], k[2])
    pages: list[dict] = []
    for group_key in sorted(rides_by_group.keys(), key=sort_key):
        driver_rides = rides_by_group[group_key]
        driver_name, _, _ = group_key
        header = _group_page_header(driver_rides, driver_name)
        sheet_title_base = header.pop("sheet_title_base")
        line_dicts = [_line_dict(r) for r in driver_rides]
        for chunk_idx, chunk in enumerate(_chunked(line_dicts, PROTOCOL_LINES_PER_SHEET)):
            title = sheet_title_base if chunk_idx == 0 else f"str. {chunk_idx + 1}"
            pages.append(
                {
                    **header,
                    "snd_base": _page_snd_header(chunk),
                    "lines": chunk,
                    "sheet_title_base": title,
                }
            )

    if plan_slots:
        _append_plan_slot_pages(pages, plan_slots, vehicle_kind_label_fn=vehicle_kind_label_fn)
    return pages


def _append_plan_slot_pages(pages: list[dict], plan_slots, *, vehicle_kind_label_fn) -> None:
    from waste_docs.services.reports.protocol_plan_slots import ProtocolPlanSlot

    sort_key = lambda s: (s.driver_last_name.lower(), s.vehicle_rz_id, s.vehicle_rz.lower())
    for slot in sorted(plan_slots, key=sort_key):
        if not isinstance(slot, ProtocolPlanSlot):
            continue
        vehicle_kind = vehicle_kind_label_fn(slot.vehicle_rz_id)
        drive_date = slot.drive_date.strftime("%d.%m.%Y")
        sheet_bits = [slot.driver_last_name]
        if slot.vehicle_rz:
            sheet_bits.append(slot.vehicle_rz)
        sheet_title_base = " ".join(sheet_bits) or "Protokol"
        line_dicts = list(slot.lines)
        chunks = list(_chunked(line_dicts, PROTOCOL_LINES_PER_SHEET))
        if not chunks:
            chunks = [[]]
        for chunk_idx, chunk in enumerate(chunks):
            title = sheet_title_base if chunk_idx == 0 else f"str. {chunk_idx + 1}"
            pages.append(
                {
                    "driver_last_name": slot.driver_last_name,
                    "vehicle_kind": vehicle_kind,
                    "drive_date": drive_date,
                    "snd_base": "",
                    "vehicle_rz": slot.vehicle_rz,
                    "lines": chunk,
                    "sheet_title_base": title,
                }
            )


def generate_protocol_xlsx(
    rides: list[Pruvodka],
    template_path: Path | None = None,
    *,
    scope: str = PROTOCOL_SCOPE_SINGLE,
    selected_date: date | None = None,
    print_warnings: list[str] | None = None,
    plan_slots=None,
) -> Path:
    if load_workbook is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")
    plan_slots = plan_slots or []
    if not rides and not plan_slots:
        raise ValueError("Nejsou dostupné žádné jízdy ani sloty plánu pro protokol.")

    template = resolve_protocol_template_path(template_path)
    if not template.is_file():
        raise FileNotFoundError(f"Šablona protokolu neexistuje: {template}")
    keep_vba = should_keep_vba_for_template(template)
    try:
        template_source = ProtocolTemplateSource(template, keep_vba=keep_vba)
    except Exception as exc:
        raise RuntimeError(
            f"Šablona protokolu nelze načíst ({template}). "
            "Makro ukládejte do souboru jizda.xlsm (Excel: Uložit jako → Sešit s podporou maker). "
            "Poškozenou šablonu jizda.xlsx obnovte ze zálohy bez maker."
        ) from exc
    try:
        template_sheet = template_source.template_sheet
        template_tokens = _collect_template_tokens(template_sheet)
        workbook = open_protocol_output_workbook(template_source)
        template_sheet_in_wb = workbook.active
        template_sheet_in_wb.title = "__template__"

        pages = build_protocol_pages(rides, plan_slots, vehicle_kind_label_fn=_vehicle_kind_label)

        if not pages:
            raise ValueError("Protokol neobsahuje žádný list k vygenerování.")

        used_titles: set[str] = set()
        for page in pages:
            sheet = template_source.clone_sheet(
                workbook,
                copy_images_fn=_copy_template_images,
            )
            sheet.title = _safe_sheet_title(page["sheet_title_base"], used_titles)
            lines = list(page.get("lines") or [])
            _render_sheet_tokens(sheet, page, template_tokens, template_sheet)
            fill_protocol_sheet(sheet, template_sheet=template_sheet, lines=lines)
            apply_poznamka_highlight(sheet, template_sheet=template_sheet)

        if "__template__" in workbook.sheetnames:
            del workbook["__template__"]

        append_protocol_notes_sheet(workbook, rides, template_sheet=template_sheet)
    finally:
        template_source.close()

    target_dir = settings.MEDIA_ROOT / "protokoly"
    target_dir.mkdir(parents=True, exist_ok=True)
    token = protocol_filename_token(rides, scope=scope, selected_date=selected_date)
    output_path = target_dir / f"protokol_{token}.xlsx"
    if output_path.exists() and scope == PROTOCOL_SCOPE_SINGLE and len(rides) == 1:
        output_path = target_dir / f"protokol_{token}_{uuid4().hex[:8]}.xlsx"
    workbook.save(output_path)
    simplex_warning = ensure_protocol_simplex_print_settings(output_path, template)
    if simplex_warning and print_warnings is not None:
        print_warnings.append(simplex_warning)
    validate_generated_xlsx(output_path)
    return output_path
