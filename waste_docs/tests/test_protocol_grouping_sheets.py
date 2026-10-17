from __future__ import annotations

import tempfile
from datetime import date, datetime, timezone as dt_timezone
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.test import TestCase, override_settings

from waste_docs.models import Pruvodka
from waste_docs.services.ride_protocol_xlsx import (
    PROTOCOL_LINES_PER_SHEET,
    generate_protocol_xlsx,
)


from waste_docs.services.reports.protocol_poznamka_highlight import (
    POZNAMKA_FILL_RGB,
    _cell_fill_rgb,
)


def _font_color_rgb(cell) -> str:
    font = getattr(cell, "font", None)
    if font is None or font.color is None:
        return ""
    rgb = getattr(font.color, "rgb", None)
    if not rgb:
        return ""
    text = str(rgb).upper()
    if len(text) == 8:
        return text[2:] if text.startswith("00") else text[-6:]
    return text


class ProtocolPoznamkaHighlightTests(TestCase):
    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_poznamka_cell_highlighted_when_note_present(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            ride = Pruvodka(
                pk=1,
                snd_cislo_listu="0000295-2026-1",
                drive_date=date(2026, 6, 10),
                vehicle_rz_id=5,
                vehicle_rz_snapshot="5SB5766",
                driver_snapshot="Jan Vrabec",
                contact_note="Rohák",
            )
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx([ride], scope="day", selected_date=date(2026, 6, 10))

            from openpyxl import load_workbook

            wb = load_workbook(path)
            sheet = wb[wb.sheetnames[0]]
            self.assertIn("Rohák", sheet["C8"].value or "")
            self.assertEqual(_cell_fill_rgb(sheet["C8"]), POZNAMKA_FILL_RGB)
            self.assertEqual(_font_color_rgb(sheet["C8"]), "000000")
            self.assertEqual(_cell_fill_rgb(sheet["C13"]), "")

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_poznamka_cell_unchanged_when_note_empty(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            ride = Pruvodka(
                pk=2,
                snd_cislo_listu="0000296-2026-1",
                drive_date=date(2026, 6, 10),
                vehicle_rz_id=5,
                vehicle_rz_snapshot="5SB5766",
                driver_snapshot="Jan Vrabec",
                contact_note="",
            )
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx([ride], scope="day", selected_date=date(2026, 6, 10))

            from openpyxl import load_workbook

            tpl = load_workbook(template_path).active
            wb = load_workbook(path)
            sheet = wb[wb.sheetnames[0]]
            self.assertEqual(sheet["C8"].value in (None, ""), True)
            self.assertEqual(_cell_fill_rgb(sheet["C8"]), _cell_fill_rgb(tpl["C8"]))


class ProtocolGroupingSheetsTests(TestCase):
    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_same_rz_driver_one_workbook_multiple_chunks(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            rides = []
            for i in range(8):
                rides.append(
                    Pruvodka(
                        pk=100 + i,
                        snd_cislo_listu=f"0000295-2026-{i + 1}",
                        drive_date=date(2026, 6, 10),
                        vehicle_rz_id=5,
                        vehicle_rz_snapshot="5SB5766",
                        driver_snapshot="Jan Vrabec",
                        created_at=datetime(2026, 6, 10, 8 + i, 0, tzinfo=dt_timezone.utc),
                    )
                )
            for r in rides:
                r.driver_id = None
                r.driver = None

            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx(rides, scope="day", selected_date=date(2026, 6, 10))

            from openpyxl import load_workbook

            wb = load_workbook(path)
            data_sheets = [s for s in wb.sheetnames if s != "Poznamky"]
            self.assertEqual(len(data_sheets), 2)
            self.assertEqual(PROTOCOL_LINES_PER_SHEET, 7)

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_different_snd_same_rz_driver_single_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            rides = [
                Pruvodka(
                    pk=1,
                    snd_cislo_listu="0000295-2026-1",
                    drive_date=date(2026, 6, 10),
                    vehicle_rz_id=5,
                    vehicle_rz_snapshot="5SB5766",
                    driver_snapshot="Jan Vrabec",
                ),
                Pruvodka(
                    pk=2,
                    snd_cislo_listu="0000296-2026-1",
                    drive_date=date(2026, 6, 10),
                    vehicle_rz_id=5,
                    vehicle_rz_snapshot="5SB5766",
                    driver_snapshot="Jan Vrabec",
                ),
            ]
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx(rides, scope="day", selected_date=date(2026, 6, 10))

            from openpyxl import load_workbook

            wb = load_workbook(path)
            data_sheets = [s for s in wb.sheetnames if s != "Poznamky"]
            self.assertEqual(len(data_sheets), 1)

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_poznamka_row_height_stays_template_size(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            ride = Pruvodka(
                pk=1,
                snd_cislo_listu="0000295-2026-1",
                drive_date=date(2026, 6, 10),
                vehicle_rz_id=5,
                vehicle_rz_snapshot="5SB5766",
                driver_snapshot="Jan Vrabec",
                contact_note="Rohák",
            )
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx([ride], scope="day", selected_date=date(2026, 6, 10))

            from openpyxl import load_workbook

            tpl = load_workbook(template_path).active
            wb = load_workbook(path)
            sheet = wb[wb.sheetnames[0]]
            self.assertIn("Rohák", sheet["C8"].value or "")
            for row_num in (8, 13, 18, 23, 28, 33, 38):
                expected = tpl.row_dimensions[row_num].height
                self.assertEqual(sheet.row_dimensions[row_num].height, expected)
                self.assertFalse(sheet.row_dimensions[row_num].hidden)
