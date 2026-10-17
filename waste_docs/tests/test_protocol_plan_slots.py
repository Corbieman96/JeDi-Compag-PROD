from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings

from catalogs.models import (
    DriverProfile,
    NameCatalog,
    NameSubjectProfile,
    NameSubjectWasteSuffixPair,
    SuffixCatalog,
    VehicleKindCatalog,
    VehicleRzCatalog,
    WasteTypeCatalog,
)
from catalogs.models.vehicle_rz import VehicleRzProfile

from waste_docs.models import SvozTask
from waste_docs.services.reports.protocol_plan_slots import plan_only_protocol_slots
from waste_docs.services.ride_bulk_actions import append_day_protocol_file
from waste_docs.services.ride_protocol_xlsx import PROTOCOL_SCOPE_DAY, generate_protocol_xlsx
from waste_docs.services.svoz_dispatch.plan import replace_yard_driver_plan_for_date


class ProtocolPlanSlotsTests(TestCase):
    databases = {"default", "addresses"}

    def setUp(self):
        self.plan_date = date(2026, 6, 20)
        User = get_user_model()
        self.user = User.objects.create_user("plan_proto", password="pw", is_superuser=True)

        nm = NameCatalog.objects.create(value="Plan proto yard")
        self.subject = NameSubjectProfile.objects.create(name=nm, company_name="CZS proto SD")
        self.subject.yard_settings.is_collection_yard = True
        self.subject.yard_settings.save()

        self.wt = WasteTypeCatalog.objects.create(code="200301", value="Objemný odpad")
        self.suf = SuffixCatalog.objects.create(value="Matrace")
        self.vk = VehicleKindCatalog.objects.create(name="Tříosý Hák", shortcut="3H", sort_order=0)
        self.driver = DriverProfile.objects.create(personal_number="D-PP", last_name="Mrňavý", first_name="Aleš")
        self.rz = VehicleRzCatalog.objects.create(value="4SJ5360")
        prof = VehicleRzProfile.objects.create(rz=self.rz)
        prof.vehicle_kinds.add(self.vk)
        NameSubjectWasteSuffixPair.objects.create(
            subject_profile=self.subject,
            waste_type=self.wt,
            suffix=self.suf,
            yard_svoz_enabled=True,
            yard_vehicle_kind=self.vk,
        )
        replace_yard_driver_plan_for_date(
            plan_date=self.plan_date,
            standard_rows=[(self.subject.pk, self.vk.pk, [(self.driver.pk, self.rz.pk, "", 1)])],
        )

    def test_plan_only_slot_without_pruvodka(self):
        slots = plan_only_protocol_slots(self.plan_date)
        self.assertEqual(len(slots), 1)
        slot = slots[0]
        self.assertEqual(slot.driver_id, self.driver.pk)
        self.assertEqual(slot.vehicle_rz_id, self.rz.pk)
        self.assertEqual(slot.lines, ())

    def test_plan_only_excludes_rz_driver_with_pruvodka(self):
        with patch(
            "waste_docs.services.reports.protocol_plan_slots.org_occupied_rz_driver_pairs",
            return_value=frozenset({(self.rz.pk, self.driver.pk)}),
        ):
            self.assertEqual(plan_only_protocol_slots(self.plan_date), [])

    def test_open_svoz_task_lines_in_slot(self):
        SvozTask.objects.create(
            subject_profile=self.subject,
            created_by=self.user,
            waste_type=self.wt,
            suffix=self.suf,
            vehicle_kind=self.vk,
            target_date=self.plan_date,
            assigned_driver=self.driver,
            assigned_vehicle_rz=self.rz,
            status=SvozTask.Status.OPEN,
        )
        slots = plan_only_protocol_slots(self.plan_date)
        self.assertEqual(len(slots), 1)
        self.assertEqual(len(slots[0].lines), 1)
        self.assertIn("Matrace", slots[0].lines[0].get("suffix_snapshot", ""))

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_generate_protocol_xlsx_plan_only(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        slots = plan_only_protocol_slots(self.plan_date)
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="3H",
                ):
                    path = generate_protocol_xlsx(
                        [],
                        scope=PROTOCOL_SCOPE_DAY,
                        selected_date=self.plan_date,
                        plan_slots=slots,
                    )
            self.assertTrue(path.is_file())
            self.assertIn(f"den_{self.plan_date.isoformat()}", path.name)

            from openpyxl import load_workbook

            wb = load_workbook(path)
            data_sheets = [s for s in wb.sheetnames if s != "Poznamky"]
            self.assertEqual(len(data_sheets), 1)

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_append_day_protocol_without_rides(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="3H",
                ):
                    rows = append_day_protocol_file(
                        [],
                        [],
                        self.plan_date,
                        force_regenerate=True,
                    )
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["kind"], "protocol")
