from __future__ import annotations

import tempfile
import zipfile
from pathlib import Path

from django.test import TestCase
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from waste_docs.services.reports.protocol_pdf_fonts import (
    FALLBACK_FONT,
    normalize_worksheet_fonts,
    normalize_xlsx_fonts,
)


class NormalizeWorksheetFontsTests(TestCase):
    def test_replaces_non_portable_fonts_and_keeps_style(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"].value = "Velikost"
        ws["A1"].font = Font(name="Play", bold=True, size=12.0)
        ws["A2"].value = "Kdo"
        ws["A2"].font = Font(name="Aptos", bold=True, size=11.0)
        ws["A3"].value = "SND"
        ws["A3"].font = Font(name="Aptos Black", bold=True, size=14.0)

        changed = normalize_worksheet_fonts(ws)

        self.assertEqual(changed, 3)
        for coord, size in (("A1", 12.0), ("A2", 11.0), ("A3", 14.0)):
            self.assertEqual(ws[coord].font.name, FALLBACK_FONT)
            self.assertTrue(ws[coord].font.bold)
            self.assertEqual(ws[coord].font.size, size)

    def test_keeps_installed_fonts(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"].value = "x"
        ws["A1"].font = Font(name="Arial", size=10.0)
        ws["A2"].value = "y"
        ws["A2"].font = Font(name="Calibri", size=11.0)

        changed = normalize_worksheet_fonts(ws)

        self.assertEqual(changed, 0)
        self.assertEqual(ws["A1"].font.name, "Arial")
        self.assertEqual(ws["A2"].font.name, "Calibri")


class NormalizeXlsxFontsTests(TestCase):
    def _build_xlsx(self, path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws["A1"].value = "Velikost"
        ws["A1"].font = Font(name="Play", bold=True, size=12.0)
        wb.save(path)

    def test_rewrites_styles_xml_to_arial(self):
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = Path(tmp) / "jizda.xlsx"
            self._build_xlsx(xlsx)

            with zipfile.ZipFile(xlsx) as archive:
                self.assertIn("Play", archive.read("xl/styles.xml").decode("utf-8"))

            changed = normalize_xlsx_fonts(xlsx)
            self.assertGreaterEqual(changed, 1)

            with zipfile.ZipFile(xlsx) as archive:
                styles = archive.read("xl/styles.xml").decode("utf-8")
            self.assertNotIn('val="Play"', styles)
            self.assertIn(FALLBACK_FONT, styles)

            wb = load_workbook(xlsx)
            self.assertEqual(wb.active["A1"].font.name, FALLBACK_FONT)

    def test_no_change_when_already_arial(self):
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = Path(tmp) / "ok.xlsx"
            wb = Workbook()
            wb.active["A1"].value = "x"
            wb.active["A1"].font = Font(name="Arial")
            wb.save(xlsx)

            self.assertEqual(normalize_xlsx_fonts(xlsx), 0)
