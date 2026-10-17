from __future__ import annotations

import os
import tempfile
import time
from datetime import date, datetime, timezone as dt_timezone
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.test import TestCase, override_settings

from waste_docs.models import Pruvodka
from waste_docs.services.ride_bulk_actions import append_day_protocol_file
from waste_docs.services.ride_protocol_xlsx import (
    PROTOCOL_SCOPE_DAY,
    PROTOCOL_SCOPE_GROUP,
    generate_protocol_xlsx,
    protocol_filename_token,
)


class ProtocolStaleRegenerationTests(TestCase):
    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_day_protocol_found_regardless_of_ride_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            protocol_dir = media / "protokoly"
            protocol_dir.mkdir(parents=True)
            (protocol_dir / "protokol_den_2026-03-01.pdf").write_bytes(b"%PDF")

            ride_a = Pruvodka(
                pk=1,
                snd_cislo_listu="0000283-2026-1",
                drive_date=date(2026, 3, 1),
                created_at=datetime(2026, 3, 1, 12, 0, tzinfo=dt_timezone.utc),
            )
            ride_b = Pruvodka(
                pk=2,
                snd_cislo_listu="0000280-2026",
                drive_date=date(2026, 3, 1),
                created_at=datetime(2026, 3, 1, 9, 0, tzinfo=dt_timezone.utc),
            )

            with override_settings(MEDIA_ROOT=media):
                rows = append_day_protocol_file(
                    [],
                    [ride_a, ride_b],
                    date(2026, 3, 1),
                    generate_if_missing=False,
                )

            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["kind"], "protocol")
            self.assertTrue(rows[0]["path"].endswith(".pdf"))

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_stale_day_protocol_is_not_regenerated_on_page_load(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            protocol_dir = media / "protokoly"
            protocol_dir.mkdir(parents=True)
            stale_path = protocol_dir / "protokol_den_2026-06-10.xlsx"
            stale_path.write_bytes(b"old")
            os.utime(stale_path, (1_000_000.0, 1_000_000.0))

            older = Pruvodka(
                pk=589,
                snd_cislo_listu="0000295-2026-1",
                drive_date=date(2026, 6, 10),
                vehicle_rz_id=7,
                vehicle_rz_snapshot="5SB5766",
                driver_snapshot="Oldřich Vrabec",
                created_at=datetime(2026, 6, 10, 8, 0, tzinfo=dt_timezone.utc),
            )
            newer = Pruvodka(
                pk=603,
                snd_cislo_listu="0000296-2026-6",
                drive_date=date(2026, 6, 10),
                vehicle_rz_id=7,
                vehicle_rz_snapshot="5SB5766",
                driver_snapshot="Oldřich Vrabec",
                created_at=datetime(2026, 6, 10, 12, 0, tzinfo=dt_timezone.utc),
            )

            with override_settings(MEDIA_ROOT=media):
                rows = append_day_protocol_file(
                    [],
                    [older, newer],
                    date(2026, 6, 10),
                    generate_if_missing=False,
                )

            self.assertEqual(len(rows), 1)
            self.assertEqual(stale_path.read_bytes(), b"old")
            self.assertTrue(rows[0].get("warning"))

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_stale_day_protocol_is_regenerated_when_force_regenerate(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            protocol_dir = media / "protokoly"
            protocol_dir.mkdir(parents=True)
            stale_path = protocol_dir / "protokol_den_2026-06-10.xlsx"
            stale_path.write_bytes(b"old")
            os.utime(stale_path, (1_000_000.0, 1_000_000.0))

            older = Pruvodka(
                pk=589,
                snd_cislo_listu="0000295-2026-1",
                drive_date=date(2026, 6, 10),
                vehicle_rz_id=7,
                vehicle_rz_snapshot="5SB5766",
                driver_snapshot="Oldřich Vrabec",
                created_at=datetime(2026, 6, 10, 8, 0, tzinfo=dt_timezone.utc),
            )
            newer = Pruvodka(
                pk=603,
                snd_cislo_listu="0000296-2026-6",
                drive_date=date(2026, 6, 10),
                vehicle_rz_id=7,
                vehicle_rz_snapshot="5SB5766",
                driver_snapshot="Oldřich Vrabec",
                created_at=datetime(2026, 6, 10, 12, 0, tzinfo=dt_timezone.utc),
            )

            with override_settings(MEDIA_ROOT=media):
                rows = append_day_protocol_file(
                    [],
                    [older, newer],
                    date(2026, 6, 10),
                    force_regenerate=True,
                )

            self.assertEqual(len(rows), 1)
            self.assertNotEqual(stale_path.read_bytes(), b"old")

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_stale_day_protocol_is_regenerated_when_regenerate_if_stale(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            protocol_dir = media / "protokoly"
            protocol_dir.mkdir(parents=True)
            stale_path = protocol_dir / "protokol_den_2026-06-10.xlsx"
            stale_path.write_bytes(b"old")
            os.utime(stale_path, (1_000_000.0, 1_000_000.0))

            older = Pruvodka(
                pk=589,
                snd_cislo_listu="0000295-2026-1",
                drive_date=date(2026, 6, 10),
                vehicle_rz_id=7,
                vehicle_rz_snapshot="5SB5766",
                driver_snapshot="Oldřich Vrabec",
                created_at=datetime(2026, 6, 10, 8, 0, tzinfo=dt_timezone.utc),
            )
            newer = Pruvodka(
                pk=603,
                snd_cislo_listu="0000296-2026-6",
                drive_date=date(2026, 6, 10),
                vehicle_rz_id=7,
                vehicle_rz_snapshot="5SB5766",
                driver_snapshot="Oldřich Vrabec",
                created_at=datetime(2026, 6, 10, 12, 0, tzinfo=dt_timezone.utc),
            )

            with override_settings(MEDIA_ROOT=media):
                with patch(
                    "waste_docs.services.ride_bulk_actions._ensure_protocol_pdf",
                    side_effect=lambda p: p,
                ), patch(
                    "waste_docs.services.ride_bulk_actions._save_protocol_note_supplement_pdf",
                    return_value=None,
                ):
                    rows = append_day_protocol_file(
                        [],
                        [older, newer],
                        date(2026, 6, 10),
                        generate_if_missing=False,
                        regenerate_if_stale=True,
                    )

            self.assertEqual(len(rows), 1)
            self.assertFalse(rows[0].get("warning"))
            self.assertNotEqual(stale_path.read_bytes(), b"old")

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_vrabec_text_is_in_regenerated_stale_day_protocol(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        galit = Pruvodka(
            pk=589,
            snd_cislo_listu="0000295-2026-1",
            drive_date=date(2026, 6, 10),
            vehicle_rz_id=6,
            vehicle_rz_snapshot="9S57595",
            driver_snapshot="Galit",
            name_snapshot="Firma Galit",
            created_at=datetime(2026, 6, 10, 8, 0, tzinfo=dt_timezone.utc),
        )
        vrabec = Pruvodka(
            pk=603,
            snd_cislo_listu="0000296-2026-6",
            drive_date=date(2026, 6, 10),
            vehicle_rz_id=7,
            vehicle_rz_snapshot="5SB5766",
            driver_snapshot="Oldřich Vrabec",
            name_snapshot="Firma Vrabec",
            created_at=datetime(2026, 6, 10, 12, 0, tzinfo=dt_timezone.utc),
        )

        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    first = generate_protocol_xlsx(
                        [galit],
                        scope=PROTOCOL_SCOPE_DAY,
                        selected_date=date(2026, 6, 10),
                    )
                os.utime(first, (1_000_000.0, 1_000_000.0))

                from openpyxl import load_workbook

                before = load_workbook(first)
                before_names = " ".join(before.sheetnames)
                before.close()
                self.assertNotIn("Vrabec", before_names)

                with patch(
                    "waste_docs.services.ride_bulk_actions._ensure_protocol_pdf",
                    side_effect=lambda p: p,
                ), patch(
                    "waste_docs.services.ride_bulk_actions._save_protocol_note_supplement_pdf",
                    return_value=None,
                ), patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    rows = append_day_protocol_file(
                        [],
                        [galit, vrabec],
                        date(2026, 6, 10),
                        generate_if_missing=False,
                        regenerate_if_stale=True,
                    )

            self.assertEqual(len(rows), 1)
            self.assertFalse(rows[0].get("warning"))
            after = load_workbook(first)
            after_text = " ".join(after.sheetnames)
            for sheet in after.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    after_text += " " + " ".join(str(v) for v in row if v is not None)
            after.close()
            self.assertIn("Vrabec", after_text)
            self.assertIn("Firma Vrabec", after_text)

    def test_group_protocol_token_is_per_vehicle(self):
        rides = [
            Pruvodka(pk=1, drive_date=date(2026, 6, 10), vehicle_rz_id=7),
            Pruvodka(pk=2, drive_date=date(2026, 6, 10), vehicle_rz_id=7),
        ]
        token = protocol_filename_token(rides, scope=PROTOCOL_SCOPE_GROUP)
        self.assertEqual(token, "skupina_2026-06-10_rz7")

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_group_protocol_overwrites_same_filename(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media):
                ride_old = Pruvodka(
                    pk=589,
                    snd_cislo_listu="0000295-2026-1",
                    drive_date=date(2026, 6, 10),
                    vehicle_rz_id=7,
                    vehicle_rz_snapshot="5SB5766",
                    driver_snapshot="Oldřich Vrabec",
                    created_at=datetime(2026, 6, 10, 8, 0, tzinfo=dt_timezone.utc),
                )
                first = generate_protocol_xlsx([ride_old], scope=PROTOCOL_SCOPE_GROUP)
                time.sleep(0.05)
                ride_new = Pruvodka(
                    pk=603,
                    snd_cislo_listu="0000296-2026-6",
                    drive_date=date(2026, 6, 10),
                    vehicle_rz_id=7,
                    vehicle_rz_snapshot="5SB5766",
                    driver_snapshot="Oldřich Vrabec",
                    created_at=datetime(2026, 6, 10, 12, 0, tzinfo=dt_timezone.utc),
                )
                second = generate_protocol_xlsx([ride_old, ride_new], scope=PROTOCOL_SCOPE_GROUP)

            self.assertEqual(first.name, second.name)
            self.assertGreaterEqual(second.stat().st_mtime, first.stat().st_mtime)
            self.assertGreater(len(second.name), 0)
