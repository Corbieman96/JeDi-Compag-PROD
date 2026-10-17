"""Protokol STASKA — skrytí řádku Kde u činnosti bez adresy plnění."""

from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.test import TestCase, override_settings

from catalogs.models import ActivityCatalog, ActivityFieldKey, ActivityFieldMode, ActivityFieldRule
from waste_docs.models import Pruvodka
from waste_docs.services.ride_numbering import apply_pruvodka_catalog_snapshots
from waste_docs.services.ride_protocol_xlsx import PROTOCOL_KDE_LABEL_COORDS, generate_protocol_xlsx


class ProtocolKdeVisibilityTests(TestCase):
    def setUp(self):
        self.activity_with_addr = ActivityCatalog.objects.create(name="Svoz s adresou")
        self.activity_no_addr = ActivityCatalog.objects.create(name="Pracovní stroj")
        ActivityFieldRule.objects.create(
            activity=self.activity_no_addr,
            field_key=ActivityFieldKey.FULFILLMENT_ADDRESS,
            mode=ActivityFieldMode.OFF,
        )

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_protocol_hides_kde_when_fulfillment_address_off(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        from openpyxl import load_workbook

        ride = Pruvodka(
            pk=1,
            activity_id=self.activity_no_addr.pk,
            snd_cislo_listu="0000400-2026-1",
            drive_date=date(2026, 6, 11),
            vehicle_rz_id=5,
            vehicle_rz_snapshot="5SB5766",
            driver_snapshot="Vrabec",
            name_snapshot="Stroj s.r.o.",
            address_snapshot="Město, 1, Ulice",
            contact_note="Staveniště u silnice",
            activity_snapshot=self.activity_no_addr.name,
        )
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Stroj",
                ):
                    path = generate_protocol_xlsx([ride], scope="day", selected_date=date(2026, 6, 11))

            sheet = load_workbook(path).active
            self.assertEqual(sheet["I5"].value in (None, ""), True)
            self.assertEqual(sheet[PROTOCOL_KDE_LABEL_COORDS[0]].value in (None, ""), True)
            self.assertIn("Staveniště u silnice", sheet["C8"].value or "")

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_protocol_shows_kde_when_fulfillment_address_required(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        from openpyxl import load_workbook

        ride = Pruvodka(
            pk=2,
            activity_id=self.activity_with_addr.pk,
            snd_cislo_listu="0000401-2026-1",
            drive_date=date(2026, 6, 11),
            vehicle_rz_id=5,
            vehicle_rz_snapshot="5SB5766",
            driver_snapshot="Novák",
            address_snapshot="Praha, 10, Hlavní",
            activity_snapshot=self.activity_with_addr.name,
        )
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx([ride], scope="day", selected_date=date(2026, 6, 11))

            sheet = load_workbook(path).active
            self.assertEqual(sheet["I5"].value, "Praha, 10, Hlavní")
            self.assertEqual(sheet[PROTOCOL_KDE_LABEL_COORDS[0]].value, "Kde")

    def test_snapshot_clears_address_when_activity_hides_fulfillment(self):
        from catalogs.models import AddressCatalog

        addr = AddressCatalog.objects.create(value="Praha, 1, Test")
        ride = Pruvodka(
            activity_id=self.activity_no_addr.pk,
            address_id=addr.pk,
            address_snapshot="",
        )
        apply_pruvodka_catalog_snapshots(ride)
        self.assertEqual(ride.address_snapshot, "")
