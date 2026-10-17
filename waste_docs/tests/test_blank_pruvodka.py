from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse

from catalogs.models import ActivityCatalog, DriverProfile, VehicleRzCatalog
from users.constants import ROLE_ADMIN
from users.services.roles import set_user_role
from waste_docs.models import Pruvodka
from waste_docs.services.blank_pruvodka import (
    archive_blank_pruvodka,
    archive_unused_blank_for_day,
    create_blank_pruvodky,
    is_unused_blank_pruvodka,
    maybe_clear_blank_flag,
)
from waste_docs.services.blank_pruvodka.constants import (
    BLANK_ACTIVITY_NAME,
    BLANK_ADDRESS_VALUE,
    BLANK_ICO_VALUE,
    BLANK_NAME_VALUE,
    BLANK_WASTE_TYPE_VALUE,
    get_blank_activity,
)
from waste_docs.services.blank_pruvodka.print_bundle import (
    build_blank_print_file_rows,
    parse_blank_print_pks,
)

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # pragma: no cover
    load_workbook = None


def _xlsx_cell_texts(path) -> list[str]:
    if load_workbook is None:
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    texts: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    texts.append(str(cell.value))
    workbook.close()
    return texts


class BlankPruvodkaTests(TestCase):
    databases = {"default", "addresses"}

    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user("blank_admin", password="pw", is_superuser=True)
        set_user_role(self.admin, ROLE_ADMIN)
        self.day = date(2026, 6, 8)
        get_blank_activity()

    def test_create_assigns_snd_without_rz(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media):
                created = create_blank_pruvodky(drive_date=self.day, count=3, user=self.admin)
        self.assertEqual(len(created), 3)
        snd_values = {r.snd_cislo_listu for r in created}
        self.assertEqual(len(snd_values), 3)
        for ride in created:
            ride.refresh_from_db()
            self.assertTrue(ride.is_blank_pruvodka)
            self.assertIsNone(ride.vehicle_rz_id)
            self.assertEqual(ride.activity.name, BLANK_ACTIVITY_NAME)

    def test_excel_contains_snd_and_date_only(self):
        if load_workbook is None or not (settings.BASE_DIR / "pruvodka.xlsx").is_file():
            self.skipTest("openpyxl nebo pruvodka.xlsx není k dispozici")
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media):
                created = create_blank_pruvodky(drive_date=self.day, count=1, user=self.admin)
                ride = created[0]
                ride.refresh_from_db()
                self.assertTrue(ride.generated_file)
                path = media / ride.generated_file.name
                texts = _xlsx_cell_texts(path)
        joined = " ".join(texts)
        self.assertIn(ride.snd_cislo_listu, joined)
        self.assertIn(self.day.strftime("%d.%m.%Y"), joined)
        for forbidden in (BLANK_NAME_VALUE, BLANK_WASTE_TYPE_VALUE, BLANK_ICO_VALUE, BLANK_ADDRESS_VALUE):
            self.assertNotIn(forbidden, joined)

    def test_unused_blank_has_empty_snapshots(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media):
                ride = create_blank_pruvodky(drive_date=self.day, count=1, user=self.admin)[0]
                ride.refresh_from_db()
        self.assertTrue(is_unused_blank_pruvodka(ride))
        self.assertEqual(ride.name_snapshot, "")
        self.assertEqual(ride.waste_type_snapshot, "")
        self.assertEqual(ride.ico_snapshot, "")

    def test_edit_clears_blank_flag(self):
        ride = create_blank_pruvodky(drive_date=self.day, count=1, user=self.admin)[0]
        other = ActivityCatalog.objects.create(name="Svoz test blank")
        ride.activity = other
        ride.vehicle_rz = VehicleRzCatalog.objects.create(value="1B99999")
        maybe_clear_blank_flag(ride)
        self.assertFalse(ride.is_blank_pruvodka)

    def test_archive_blank_single(self):
        ride = create_blank_pruvodky(drive_date=self.day, count=1, user=self.admin)[0]
        snd = ride.snd_cislo_listu
        archive_blank_pruvodka(pruvodka=ride, user=self.admin)
        ride.refresh_from_db()
        self.assertTrue(ride.is_archived)
        self.assertEqual(ride.snd_cislo_listu, snd)
        self.assertFalse(
            Pruvodka.objects.filter(drive_date=self.day, is_archived=False, is_blank_pruvodka=True).exists()
        )

    def test_archive_blank_bulk(self):
        create_blank_pruvodky(drive_date=self.day, count=4, user=self.admin)
        archived = archive_unused_blank_for_day(drive_date=self.day, user=self.admin)
        self.assertEqual(archived, 4)

    def test_archive_rejects_filled_blank(self):
        ride = create_blank_pruvodky(drive_date=self.day, count=1, user=self.admin)[0]
        ride.driver = DriverProfile.objects.create(personal_number="D001", first_name="A", last_name="B")
        ride.save()
        self.assertFalse(is_unused_blank_pruvodka(ride))

    def test_create_view_redirects_to_print(self):
        self.client.force_login(self.admin)
        url = reverse("waste_docs:blank_pruvodka_create")
        response = self.client.post(
            url,
            {"drive_date": self.day.isoformat(), "count": "2", "note": "test"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertIn("bianko/print", response.url)
        self.assertIn("pks=", response.url)

    def test_parse_blank_print_pks_accepts_pk_alias(self):
        self.assertEqual(parse_blank_print_pks(pks_raw="1,2"), [1, 2])
        self.assertEqual(parse_blank_print_pks(pk_raw="42"), [42])
        self.assertEqual(parse_blank_print_pks(), [])

    def test_blank_print_view_accepts_pk_param(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media):
                created = create_blank_pruvodky(drive_date=self.day, count=1, user=self.admin)
        self.client.force_login(self.admin)
        url = reverse("waste_docs:blank_pruvodka_print")
        response = self.client.get(f"{url}?pk={created[0].pk}")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Vytisknout jedním PDF")

    def test_build_blank_print_file_rows_no_protocol(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media):
                created = create_blank_pruvodky(drive_date=self.day, count=2, user=self.admin)
                pks = [r.pk for r in created]
                _, file_rows, _, prepare_ctx = build_blank_print_file_rows(pks, self.admin)
        self.assertGreaterEqual(len(file_rows), 2)
        self.assertFalse(prepare_ctx["needs_print_prepare"])
        kinds = {row.get("kind") for row in file_rows}
        self.assertEqual(kinds, {"xlsx"})
