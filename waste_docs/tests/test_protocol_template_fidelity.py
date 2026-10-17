from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.test import TestCase, override_settings

from waste_docs.models import Pruvodka
from waste_docs.services.reports.protocol_sheet_bounds import iter_protocol_cells
from waste_docs.services.ride_protocol_xlsx import (
    PROTOCOL_STATIC_LABEL_COORDS,
    generate_protocol_xlsx,
)


class ProtocolTemplateFidelityTests(TestCase):
    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_static_labels_match_jizda_template(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        from openpyxl import load_workbook

        tpl_wb = load_workbook(template_path)
        tpl_sheet = tpl_wb.active
        expected = {coord: tpl_sheet[coord].value for coord in PROTOCOL_STATIC_LABEL_COORDS}

        ride = Pruvodka(
            pk=1,
            snd_cislo_listu="0000295-2026-1",
            drive_date=date(2026, 6, 10),
            vehicle_rz_id=5,
            vehicle_rz_snapshot="5SB5766",
            driver_snapshot="Jan Vrabec",
            name_snapshot="Firma s.r.o.",
            suffix_snapshot="Spalovna",
        )
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx([ride], scope="day", selected_date=date(2026, 6, 10))

            out_wb = load_workbook(path)
            out_sheet = out_wb[out_wb.sheetnames[0]]
            for coord, value in expected.items():
                self.assertEqual(out_sheet[coord].value, value, coord)

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_token_cell_keeps_template_font(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        from openpyxl import load_workbook

        tpl_wb = load_workbook(template_path)
        tpl_font = tpl_wb.active["C2"].font

        ride = Pruvodka(
            pk=1,
            snd_cislo_listu="0000295-2026-1",
            drive_date=date(2026, 6, 10),
            vehicle_rz_id=5,
            vehicle_rz_snapshot="5SB5766",
            driver_snapshot="Jan Vrabec",
        )
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx([ride], scope="day", selected_date=date(2026, 6, 10))

            out_wb = load_workbook(path)
            out_font = out_wb[out_wb.sheetnames[0]]["C2"].font
            self.assertEqual(out_font.name, tpl_font.name)
            self.assertEqual(out_font.bold, tpl_font.bold)

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_inserted_values_preserve_mixed_case(self):
        ride = Pruvodka(
            pk=1,
            snd_cislo_listu="0000295-2026-1",
            drive_date=date(2026, 6, 10),
            vehicle_rz_id=5,
            vehicle_rz_snapshot="5SB5766",
            driver_snapshot="Jan Vrabec",
            name_snapshot="Firma s.r.o.",
            suffix_snapshot="Spalovna",
            contact_note="Poznámka s Českými písmeny",
        )
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx([ride], scope="day", selected_date=date(2026, 6, 10))

            from openpyxl import load_workbook

            out_wb = load_workbook(path)
            sheet = out_wb[out_wb.sheetnames[0]]
            self.assertEqual(sheet["C5"].value, "Firma s.r.o.")
            self.assertEqual(sheet["I6"].value, "Spalovna")
            self.assertIn("Poznámka s Českými písmeny", sheet["C8"].value or "")

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_all_sheet_fonts_match_jizda_template(self):
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        from openpyxl import load_workbook

        tpl_wb = load_workbook(template_path)
        tpl_sheet = tpl_wb.active
        tpl_fonts = {
            cell.coordinate: cell.font.name
            for cell in iter_protocol_cells(tpl_sheet)
            if cell.font and cell.font.name
        }

        ride = Pruvodka(
            pk=1,
            snd_cislo_listu="0000295-2026-1",
            drive_date=date(2026, 6, 10),
            vehicle_rz_id=5,
            vehicle_rz_snapshot="5SB5766",
            driver_snapshot="Jan Vrabec",
            name_snapshot="Firma s.r.o.",
            suffix_snapshot="Spalovna",
            contact_note="Poznámka s Českými písmeny",
        )
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx([ride], scope="day", selected_date=date(2026, 6, 10))

            out_wb = load_workbook(path)
            out_sheet = out_wb[out_wb.sheetnames[0]]
            mismatches = []
            for coord, expected_name in tpl_fonts.items():
                actual_name = out_sheet[coord].font.name if out_sheet[coord].font else None
                if actual_name != expected_name:
                    mismatches.append(f"{coord}: {expected_name!r} != {actual_name!r}")
            self.assertEqual(mismatches, [], "\n".join(mismatches[:20]))

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_template_has_no_aptos_fonts(self):
        """Aptos vzniká při uložení šablony v Excelu 2024+ — protokol má používat Play jako ve zdroji."""
        template_path = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template_path.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        from openpyxl import load_workbook

        tpl_sheet = load_workbook(template_path).active
        aptos_cells = [
            cell.coordinate
            for cell in iter_protocol_cells(tpl_sheet)
            if cell.font and cell.font.name and "Aptos" in cell.font.name
        ]
        self.assertEqual(aptos_cells, [], f"Aptos ve šabloně: {aptos_cells[:15]}")
