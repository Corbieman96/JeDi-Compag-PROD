from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.test import TestCase, override_settings

from waste_docs.models import Pruvodka
from waste_docs.services.reports.protocol_page_setup import (
    PROTOCOL_SHEET_PRINT_AREA,
    PROTOCOL_SHEET_PRINT_AREA_BACK,
    PROTOCOL_SHEET_PRINT_AREA_FRONT,
    PROTOCOL_SHEET_PRINT_AREA_FULL,
    PROTOCOL_SHEET_PRINT_AREAS,
    apply_protocol_page_setup_from_template,
    apply_single_side_protocol_page_setup,
    is_protocol_xlsx_path,
)
from waste_docs.services.reports.xlsx_to_pdf import _find_soffice, convert_xlsx_to_pdf
from waste_docs.services.ride_protocol_xlsx import generate_protocol_xlsx

try:
    from openpyxl import load_workbook
    from pypdf import PdfReader
except ImportError:
    load_workbook = None
    PdfReader = None


class ProtocolPageSetupTests(TestCase):
    def test_apply_single_side_page_setup(self):
        if load_workbook is None:
            self.skipTest("openpyxl není nainstalovaný")
        template = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        wb = load_workbook(template)
        tpl = wb.active
        ws = wb.copy_worksheet(tpl)
        apply_single_side_protocol_page_setup(ws, tpl)
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(ws.page_setup.orientation, tpl.page_setup.orientation)
        area = ws.print_area.replace("'", "").replace("$", "")
        self.assertIn("A1:M38", area)
        self.assertNotIn("A40:M79", area)
        self.assertEqual(ws.row_breaks.brk, [])
        wb.close()

    def test_apply_worksheet_sets_dual_print_areas_from_template(self):
        if load_workbook is None:
            self.skipTest("openpyxl není nainstalovaný")
        template = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        wb = load_workbook(template)
        tpl = wb.active
        ws = wb.copy_worksheet(tpl)
        apply_protocol_page_setup_from_template(ws, tpl)
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(ws.page_setup.fitToHeight, tpl.page_setup.fitToHeight)
        self.assertEqual(ws.page_setup.orientation, "landscape")
        area = ws.print_area.replace("'", "")
        self.assertIn("$A$1:$M$39", area)
        self.assertIn("$A$40:$M$79", area)
        self.assertEqual(ws.row_breaks.brk, [])
        wb.close()

    def test_print_areas_constants(self):
        self.assertEqual(PROTOCOL_SHEET_PRINT_AREA, "A1:M38")
        self.assertEqual(
            PROTOCOL_SHEET_PRINT_AREAS,
            (PROTOCOL_SHEET_PRINT_AREA_FRONT, PROTOCOL_SHEET_PRINT_AREA_BACK),
        )
        self.assertEqual(PROTOCOL_SHEET_PRINT_AREA_FRONT, "A1:M39")
        self.assertEqual(PROTOCOL_SHEET_PRINT_AREA_BACK, "A40:M79")
        self.assertEqual(PROTOCOL_SHEET_PRINT_AREA_FULL, "A1:M79")

    def test_generate_protocol_inherits_template_page_setup(self):
        if load_workbook is None:
            self.skipTest("openpyxl není nainstalovaný")
        base_template = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not base_template.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        tpl_wb = load_workbook(base_template)
        tpl = tpl_wb.active

        ride = Pruvodka(
            pk=1,
            snd_cislo_listu="0000295-2026-1",
            drive_date=date(2026, 6, 10),
            vehicle_rz_id=5,
            vehicle_rz_snapshot="5SB5766",
            driver_snapshot="Jan Vrabec",
        )
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx(
                        [ride],
                        scope="day",
                        selected_date=date(2026, 6, 10),
                    )
                self.assertTrue(is_protocol_xlsx_path(path))
                wb = load_workbook(path)
                sheet = wb[wb.sheetnames[0]]
                self.assertEqual(sheet.page_setup.fitToHeight, tpl.page_setup.fitToHeight)
                self.assertEqual(sheet.page_setup.fitToWidth, tpl.page_setup.fitToWidth)
                self.assertEqual(sheet.page_setup.orientation, tpl.page_setup.orientation)
                tpl_area = (tpl.print_area or "").replace("'", "").replace("$", "")
                out_area = (sheet.print_area or "").replace("'", "").replace("$", "")
                if tpl_area:
                    self.assertEqual(out_area, tpl_area)
                self.assertEqual(sheet.row_breaks.brk, tpl.row_breaks.brk)
                wb.close()
        tpl_wb.close()

    def test_template_page_setup_on_copy(self):
        if load_workbook is None:
            self.skipTest("openpyxl není nainstalovaný")
        template = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not template.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        with tempfile.TemporaryDirectory() as tmp:
            copy_path = Path(tmp) / "jizda_copy.xlsx"
            copy_path.write_bytes(template.read_bytes())
            from waste_docs.services.reports.protocol_sheet_bounds import trim_protocol_sheet_to_bounds

            wb = load_workbook(copy_path)
            ws = wb.active
            trim_protocol_sheet_to_bounds(ws)
            apply_single_side_protocol_page_setup(ws, load_workbook(template).active)
            wb.save(copy_path)
            wb.close()

            wb = load_workbook(copy_path)
            ws = wb.active
            tpl = load_workbook(template).active
            self.assertEqual(ws.page_setup.fitToHeight, tpl.page_setup.fitToHeight)
            area = ws.print_area.replace("'", "")
            self.assertIn("A1:M38", area.replace("$", ""))
            self.assertEqual(ws.row_breaks.brk, [])
            self.assertEqual(ws.max_row, 38)
            wb.close()

    def test_protocol_convert_one_page_per_sheet(self):
        if load_workbook is None or PdfReader is None:
            self.skipTest("openpyxl/pypdf není nainstalovaný")
        if _find_soffice() is None:
            self.skipTest("LibreOffice není k dispozici")

        base_template = Path(settings.BASE_DIR) / "jizda.xlsx"
        if not base_template.is_file():
            self.skipTest("jizda.xlsx není v kořeni projektu")

        ride = Pruvodka(
            pk=1,
            snd_cislo_listu="0000396-2026-1",
            drive_date=date(2026, 6, 18),
            vehicle_rz_id=1,
            vehicle_rz_snapshot="S070164",
            driver_snapshot="Cizi",
        )
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media, BASE_DIR=settings.BASE_DIR):
                with patch(
                    "waste_docs.services.ride_protocol_xlsx._vehicle_kind_label",
                    return_value="Solo",
                ):
                    path = generate_protocol_xlsx(
                        [ride],
                        scope="single",
                    )
                    pdf_path = convert_xlsx_to_pdf(path)
                self.assertEqual(len(PdfReader(str(pdf_path)).pages), 1)
