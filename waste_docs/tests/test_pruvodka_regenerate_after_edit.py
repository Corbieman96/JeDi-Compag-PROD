from __future__ import annotations

import os
import tempfile
from datetime import date, datetime, timezone as dt_timezone
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings

from catalogs.models import (
    AddressCatalog,
    IcoCatalog,
    NameCatalog,
    SuffixCatalog,
    VehicleRzCatalog,
    WasteCategoryCatalog,
    WasteCodeCatalog,
    WasteTypeCatalog,
)
from users.constants import ROLE_ADMIN
from users.services.roles import set_user_role
from waste_docs.models import Pruvodka
from waste_docs.services.pruvodka_generated_files import (
    ensure_excel_files_for_rides,
    regenerate_all_excels,
    regenerate_pruvodka_excel,
)
from waste_docs.services.ride_protocol_xlsx import protocol_file_is_stale

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # pragma: no cover
    load_workbook = None


def _xlsx_cell_texts(path: Path) -> list[str]:
    if load_workbook is None:
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    texts: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    texts.append(str(cell.value))
    workbook.close()
    return texts


class PruvodkaRegenerateAfterEditTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user("regen_admin", password="pw")
        set_user_role(self.admin, ROLE_ADMIN)
        self.ride_date = date(2026, 6, 11)
        name = NameCatalog.objects.create(value="Subjekt regen test")
        common = dict(
            name=name,
            ico=IcoCatalog.objects.create(value="55667788"),
            address=AddressCatalog.objects.create(value="Adr regen"),
            waste_type=WasteTypeCatalog.objects.create(code="880088", value="Wregen"),
            waste_code=WasteCodeCatalog.objects.create(value="WCRE"),
            waste_category=WasteCategoryCatalog.objects.create(value="KORE"),
            suffix=SuffixCatalog.objects.create(value="SFXRE"),
            drive_date=self.ride_date,
            vehicle_rz=VehicleRzCatalog.objects.create(value="1P55555"),
            vehicle_rz_snapshot="1P55555",
            driver_snapshot="Jan Regen",
            snd_cislo_listu="0000999-2026",
            created_by=self.admin,
        )
        self.ride = Pruvodka.objects.create(container_number="", **common)

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_generate_twice_includes_container(self):
        if load_workbook is None or not (settings.BASE_DIR / "pruvodka.xlsx").is_file():
            self.skipTest("openpyxl nebo pruvodka.xlsx není k dispozici")

        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media):
                regenerate_pruvodka_excel(self.ride)
                self.ride.container_number = "KON-12345"
                self.ride.save(update_fields=("container_number",))
                path = regenerate_pruvodka_excel(self.ride)

            self.assertIsNotNone(path)
            texts = _xlsx_cell_texts(path)
            self.assertTrue(any("KON-12345" in text for text in texts), texts)

    @patch("waste_docs.services.pruvodka_generated_files.generate_xlsx")
    def test_regenerate_all_deletes_existing_first(self, mock_generate):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            (media / "pruvodky").mkdir(parents=True)
            old_rel = "pruvodky/old_pruvodka.xlsx"
            old_path = media / old_rel
            old_path.write_bytes(b"old")
            self.ride.generated_file.name = old_rel
            self.ride.status = Pruvodka.Status.GENERATED
            self.ride.save(update_fields=("generated_file", "status"))

            new_path = media / "pruvodky" / "new_pruvodka.xlsx"
            mock_generate.return_value = str(new_path)

            with override_settings(MEDIA_ROOT=media):
                count = regenerate_all_excels([self.ride])

            self.assertEqual(count, 1)
            self.assertFalse(old_path.exists())
            mock_generate.assert_called_once()
            self.ride.refresh_from_db()
            self.assertTrue(self.ride.generated_file.name.endswith("new_pruvodka.xlsx"))

    @patch("waste_docs.services.pruvodka_generated_files.generate_xlsx")
    def test_ensure_excel_skips_existing_file(self, mock_generate):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            (media / "pruvodky").mkdir(parents=True)
            existing_rel = "pruvodky/existing_pruvodka.xlsx"
            existing_path = media / existing_rel
            existing_path.write_bytes(b"existing")
            self.ride.generated_file.name = existing_rel
            self.ride.status = Pruvodka.Status.GENERATED
            self.ride.save(update_fields=("generated_file", "status"))

            with override_settings(MEDIA_ROOT=media):
                count = ensure_excel_files_for_rides([self.ride])

            self.assertEqual(count, 0)
            mock_generate.assert_not_called()
            self.assertTrue(existing_path.exists())

    @patch("waste_docs.services.pruvodka_generated_files.generate_xlsx")
    def test_ensure_excel_generates_when_missing(self, mock_generate):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            (media / "pruvodky").mkdir(parents=True)
            new_path = media / "pruvodky" / "new_pruvodka.xlsx"
            mock_generate.return_value = str(new_path)

            with override_settings(MEDIA_ROOT=media):
                count = ensure_excel_files_for_rides([self.ride])

            self.assertEqual(count, 1)
            mock_generate.assert_called_once()

    def test_protocol_stale_after_edit(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            protocol_dir = media / "protokoly"
            protocol_dir.mkdir(parents=True)
            stale_path = protocol_dir / "protokol_den_2026-06-11.xlsx"
            stale_path.write_bytes(b"old")
            file_mtime = datetime(2026, 6, 11, 9, 0, tzinfo=dt_timezone.utc).timestamp()
            os.utime(stale_path, (file_mtime, file_mtime))

            created = datetime(2026, 6, 11, 8, 0, tzinfo=dt_timezone.utc)
            updated = datetime(2026, 6, 11, 8, 0, tzinfo=dt_timezone.utc)
            ride = Pruvodka(
                pk=self.ride.pk,
                snd_cislo_listu="0000999-2026",
                drive_date=self.ride_date,
                vehicle_rz_id=1,
                vehicle_rz_snapshot="1P55555",
                driver_snapshot="Jan Regen",
                created_at=created,
                updated_at=updated,
            )
            self.assertFalse(protocol_file_is_stale([ride], stale_path))

            ride.updated_at = datetime(2026, 6, 11, 14, 0, tzinfo=dt_timezone.utc)
            self.assertTrue(protocol_file_is_stale([ride], stale_path))
