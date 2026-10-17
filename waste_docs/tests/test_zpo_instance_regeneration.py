from __future__ import annotations

import io
import tempfile
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings

from catalogs.models import (
    ActivityCatalog,
    AddressCatalog,
    IcoCatalog,
    NameCatalog,
    NameSubjectProfile,
    NameSubjectWasteSuffixPair,
    SuffixCatalog,
    VehicleRzCatalog,
    WasteCategoryCatalog,
    WasteCodeCatalog,
    WasteTypeCatalog,
    ZpoDocument,
)
from waste_docs.models import Pruvodka
from waste_docs.services.zpo_instance_staleness import zpo_instance_needs_regeneration
from waste_docs.services.zpo_xlsx_stamp import build_zpo_xlsx_bytes
from waste_docs.zpo_instance_generator import CURRENT_ZPO_GENERATOR_REVISION, ensure_zpo_instance_for_pruvodka
from waste_docs.zpo_resolved import zpo_kind_for_file

try:
    from openpyxl import Workbook, load_workbook
except ModuleNotFoundError:  # pragma: no cover
    Workbook = None
    load_workbook = None


def _minimal_pdf_bytes() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=400, height=600)
    bio = io.BytesIO()
    writer.write(bio)
    return bio.getvalue()


class ZpoInstanceRegenerationTests(TestCase):
    def _setup_catalogs_in_media(self, media: Path):
        User = get_user_model()
        self.user = User.objects.create_user("zpo_regen", password="pw")
        self.name = NameCatalog.objects.create(value="Subjekt ZPO regen")
        self.ico = IcoCatalog.objects.create(value="11223344")
        self.profile = NameSubjectProfile.objects.create(name=self.name, ico=self.ico)
        self.wt = WasteTypeCatalog.objects.create(code="170101", value="W regen")
        self.suf = SuffixCatalog.objects.create(value="Konc regen")
        self.activity = ActivityCatalog.objects.create(name="Výměna regen")
        self.zpo_doc = ZpoDocument.objects.create(label="Vzor regen")
        self.zpo_doc.file.save("zpo/vzor_regen.pdf", ContentFile(_minimal_pdf_bytes()), save=True)
        self.pair = NameSubjectWasteSuffixPair.objects.create(
            subject_profile=self.profile,
            waste_type=self.wt,
            suffix=self.suf,
            zpo=self.zpo_doc,
        )
        self.ride = Pruvodka.objects.create(
            activity=self.activity,
            drive_date=date(2026, 7, 9),
            name=self.name,
            ico=self.ico,
            address=AddressCatalog.objects.create(value="Adr regen"),
            waste_type=self.wt,
            waste_code=WasteCodeCatalog.objects.create(value="WCREG"),
            waste_category=WasteCategoryCatalog.objects.create(value="KOREG"),
            suffix=self.suf,
            vehicle_rz=VehicleRzCatalog.objects.create(value="1P22222"),
            container_number="",
            created_by=self.user,
        )
        Pruvodka.objects.filter(pk=self.ride.pk).update(
            ride_group_number=222,
            ride_sequence=1,
            snd_cislo_listu="0000222-2026-1",
        )
        self.ride.refresh_from_db()

    def test_ensure_regenerates_when_revision_outdated(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media):
                self._setup_catalogs_in_media(media)
                instance_rel = f"pruvodky/zpo_instance/{self.ride.pk}_old.pdf"
                instance_path = media / instance_rel
                instance_path.parent.mkdir(parents=True, exist_ok=True)
                instance_path.write_bytes(b"%PDF-old-without-snd")

                self.ride.zpo_instance_file.name = instance_rel
                self.ride.zpo_instance_for_snd = self.ride.snd_cislo_listu
                self.ride.zpo_generator_revision = max(0, CURRENT_ZPO_GENERATOR_REVISION - 1)
                self.ride.save(
                    update_fields=("zpo_instance_file", "zpo_instance_for_snd", "zpo_generator_revision")
                )

                self.assertTrue(
                    zpo_instance_needs_regeneration(self.ride, pair=self.pair),
                    "zastaralá revize musí vyžadovat přegenerování",
                )
                ok = ensure_zpo_instance_for_pruvodka(self.ride)
                self.assertTrue(ok)
                self.ride.refresh_from_db()
                self.assertEqual(self.ride.zpo_generator_revision, CURRENT_ZPO_GENERATOR_REVISION)
                self.assertTrue(self.ride.zpo_instance_file.name)
                self.assertTrue((media / self.ride.zpo_instance_file.name).is_file())

                from pypdf import PdfReader

                pdf_path = media / self.ride.zpo_instance_file.name
                with pdf_path.open("rb") as pdf_file:
                    pdf_text = "".join(
                        page.extract_text() or ""
                        for page in PdfReader(pdf_file).pages
                    )
                self.assertIn("0000222-2026-1", pdf_text)

    def test_ensure_skips_when_instance_up_to_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media):
                self._setup_catalogs_in_media(media)
                instance_rel = f"pruvodky/zpo_instance/{self.ride.pk}_current.pdf"
                instance_path = media / instance_rel
                instance_path.parent.mkdir(parents=True, exist_ok=True)
                instance_path.write_bytes(_minimal_pdf_bytes())

                self.ride.zpo_instance_file.name = instance_rel
                self.ride.zpo_instance_for_snd = self.ride.snd_cislo_listu
                self.ride.zpo_generator_revision = CURRENT_ZPO_GENERATOR_REVISION
                self.ride.save(
                    update_fields=("zpo_instance_file", "zpo_instance_for_snd", "zpo_generator_revision")
                )

                self.assertFalse(zpo_instance_needs_regeneration(self.ride, pair=self.pair))
                with patch(
                    "waste_docs.zpo_instance_generator.save_zpo_instance_bytes"
                ) as mock_save:
                    ok = ensure_zpo_instance_for_pruvodka(self.ride)
                self.assertTrue(ok)
                mock_save.assert_not_called()

    def test_force_regenerates_even_when_up_to_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            media = Path(tmp)
            with override_settings(MEDIA_ROOT=media):
                self._setup_catalogs_in_media(media)
                instance_rel = f"pruvodky/zpo_instance/{self.ride.pk}_current.pdf"
                instance_path = media / instance_rel
                instance_path.parent.mkdir(parents=True, exist_ok=True)
                instance_path.write_bytes(b"%PDF-old")

                self.ride.zpo_instance_file.name = instance_rel
                self.ride.zpo_instance_for_snd = self.ride.snd_cislo_listu
                self.ride.zpo_generator_revision = CURRENT_ZPO_GENERATOR_REVISION
                self.ride.save(
                    update_fields=("zpo_instance_file", "zpo_instance_for_snd", "zpo_generator_revision")
                )

                with patch(
                    "waste_docs.zpo_instance_generator.save_zpo_instance_bytes",
                    wraps=None,
                ) as mock_save:
                    mock_save.return_value = instance_rel
                    ok = ensure_zpo_instance_for_pruvodka(self.ride, force=True)
                self.assertTrue(ok)
                mock_save.assert_called_once()


class ZpoXlsxStampTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user("zpo_xlsx", password="pw")
        self.ride = Pruvodka(
            pk=42,
            snd_cislo_listu="0000222-2026-1",
            created_by=self.user,
        )

    def test_zpo_kind_recognizes_xlsx(self):
        file_obj = MagicMock()
        file_obj.name = "zpo/template.xlsx"
        self.assertEqual(zpo_kind_for_file(file_obj), "xlsx")

    def test_build_zpo_xlsx_replaces_snd_placeholder(self):
        if Workbook is None or load_workbook is None:
            self.skipTest("openpyxl není k dispozici")

        bio = io.BytesIO()
        workbook = Workbook()
        workbook.active["A1"] = "List: {{SND_CISLO_LISTU_}}"
        workbook.save(bio)
        template_bytes = bio.getvalue()

        template = MagicMock()
        template.open.return_value.__enter__ = lambda s: io.BytesIO(template_bytes)
        template.open.return_value.__exit__ = lambda s, *a: None

        out = build_zpo_xlsx_bytes(template, self.ride)
        result_wb = load_workbook(io.BytesIO(out), read_only=True, data_only=True)
        self.assertEqual(result_wb.active["A1"].value, "List: 0000222-2026-1")
        result_wb.close()
