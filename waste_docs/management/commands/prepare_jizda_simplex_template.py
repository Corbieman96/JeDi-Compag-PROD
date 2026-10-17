from __future__ import annotations

from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError

from waste_docs.services.reports.excel_template_duplex import apply_simplex_via_excel
from waste_docs.services.reports.protocol_print_settings import (
    apply_printer_settings_blob,
    read_template_printer_settings,
    strip_printer_settings_from_xlsx,
)
from waste_docs.services.reports.protocol_sheet_bounds import trim_protocol_sheet_to_bounds
from waste_docs.services.reports.simplex_printer_blob import (
    capture_simplex_blob_for_ricoh,
    is_simplex_template_blob,
    is_trusted_simplex_printer_settings_blob,
    save_trusted_simplex_blob,
)

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None  # type: ignore[assignment,misc]


class Command(BaseCommand):
    help = (
        "Pripravi jizda.xlsx: A1:M38, landscape, jednostranny tisk. "
        "printerSettings ulozi MS Excel nebo Windows spooler (RICOH). "
        "Commitnete jizda.xlsx do gitu — server kopiruje nastaveni do protokolu."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--template",
            default="",
            help="Cesta ke sablone (vychozi: BASE_DIR/jizda.xlsx).",
        )
        parser.add_argument(
            "--printer",
            default="RICOH",
            help='Cast nazvu tiskarny pro capture pres spooler (vychozi: "RICOH").',
        )

    def handle(self, *args, **options):
        if load_workbook is None:
            raise CommandError("Knihovna openpyxl neni nainstalovana.")

        template = (
            Path(options["template"]).resolve()
            if options["template"]
            else Path(settings.BASE_DIR) / "jizda.xlsx"
        )
        if not template.is_file():
            raise CommandError(f"Sablona neexistuje: {template}")

        wb = load_workbook(template)
        trim_protocol_sheet_to_bounds(wb.active)
        wb.save(template)
        wb.close()

        strip_printer_settings_from_xlsx(template)

        try:
            apply_simplex_via_excel(template)
        except RuntimeError as exc:
            self.stdout.write(self.style.WARNING(str(exc)))

        blob = read_template_printer_settings(template)
        if not (blob and is_simplex_template_blob(blob)):
            self.stdout.write("Excel neulozil printerSettings — zkousim capture pres spooler...")
            try:
                captured = capture_simplex_blob_for_ricoh(options["printer"] or "RICOH")
                apply_printer_settings_blob(template, captured)
                try:
                    save_trusted_simplex_blob(captured)
                    self.stdout.write(
                        self.style.SUCCESS(
                            "Zaloha blobu tez ulozena do data/ricoh_im_c300_simplex.bin (volitelny fallback)."
                        )
                    )
                except Exception:
                    pass
            except RuntimeError as exc:
                self.stdout.write(self.style.WARNING(str(exc)))

        blob = read_template_printer_settings(template)
        if blob and is_trusted_simplex_printer_settings_blob(blob):
            self.stdout.write(
                self.style.SUCCESS(
                    f"{template.name}: jednostranny tisk ulozen ({len(blob)} B printerSettings). "
                    "Commitnete jizda.xlsx do gitu a nasadte na server."
                )
            )
            return

        raise CommandError(
            f"{template.name} stale nema jednostranny printerSettings. "
            "Otevřete sablonu v MS Excel na PC s RICOH, nastavte jednostranny tisk "
            "(A4 na sirku, oblast A1:M38) a ulozte, nebo nainstalujte ovladac RICOH "
            "a spustte tento prikaz znovu."
        )
