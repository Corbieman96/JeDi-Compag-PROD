import json
import re
from pathlib import Path
from uuid import uuid4

from django.db import models
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from django.conf import settings

from catalogs.models import NameSubjectProfile, SuffixSiteProfile, WasteTypeCatalog
from waste_docs.models import Pruvodka

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # pragma: no cover
    load_workbook = None


PLACEHOLDER_MAP = {
    "{{KDO_NAME}}": "name_snapshot",
    "{{KDO_ICO}}": "ico_snapshot",
    "{{KDO_ADDRESS}}": "address_snapshot",
    "{{WASTE_TYPE}}": "waste_type_snapshot",
    "{{WASTE_CODE}}": "waste_code_snapshot",
    "{{WASTE_CATEGORY}}": "waste_category_snapshot",
    "{{VEHICLE_RZ}}": "vehicle_rz_snapshot",
    "{{KONCOVKA}}": "suffix_snapshot",
    "{{SND_CISLO_LISTU_}}": "snd_cislo_listu",
    # Aliasy požadované ve šabloně pruvodka.xlsx
    "{{KONTEJNER_}}": "container_number",
    "{{KDO_}}": "name_snapshot",
    "{{DRUH_ODPADU_}}": "waste_type_snapshot",
    "{{KOD_}}": "waste_code_snapshot",
    "{{KATEGORIE_}}": "waste_category_snapshot",
    "{{MISTO_DODANI_CINNOST_}}": "address_snapshot",
    "{{RZ_VOZIDLA}}": "vehicle_rz_snapshot",
    "{{KONCOVKA_}}": "suffix_snapshot",
    "{{SMLOUVA_}}": "contract_number_snapshot",
    "{{HMOTNOST_}}": "weight_snapshot",
}


def _facility_address_value_for_pruvodka(pruvodka: Pruvodka) -> str:
    """Adresa provozovny ze subjektu (katalog name) — ne Adresa plnění."""
    name_id = getattr(pruvodka, "name_id", None)
    if not name_id:
        return ""
    prof = (
        NameSubjectProfile.objects.filter(name_id=name_id)
        .select_related("address_profile__facility_address")
        .first()
    )
    if not prof or not prof.facility_address_id or not prof.facility_address:
        return ""
    return (prof.facility_address.value or "").strip()


def _fulfillment_address_value_for_pruvodka(pruvodka: Pruvodka) -> str:
    """Adresa plnění z Evidence jízdy — ne Adresa provozovny."""
    if getattr(pruvodka, "address_id", None) and pruvodka.address:
        val = (pruvodka.address.value or "").strip()
        if val:
            return val
    return (getattr(pruvodka, "address_snapshot", None) or "").strip()


def _placeholder_values(pruvodka: Pruvodka) -> dict[str, str]:
    from waste_docs.services.blank_pruvodka.excel import blank_excel_placeholder_values, is_blank_excel_minimal

    if is_blank_excel_minimal(pruvodka):
        return blank_excel_placeholder_values(pruvodka)

    values: dict[str, str] = {}
    for token, attr in PLACEHOLDER_MAP.items():
        values[token] = getattr(pruvodka, attr, "") or ""

    # Dvě různé adresy — nikdy neprohodit:
    # {{ADRESA_}} = provozovna (katalog); {{MISTO_DODANI_CINNOST_}} = plnění (jízda).
    facility = _facility_address_value_for_pruvodka(pruvodka)
    fulfillment = _fulfillment_address_value_for_pruvodka(pruvodka)
    ico = values.get("{{KDO_ICO}}", "")
    values["{{ADRESA_}}"] = f"{facility} / {ico}".strip(" /")
    values["{{MISTO_DODANI_CINNOST_}}"] = fulfillment
    values["{{DATUM}}"] = pruvodka.drive_date.strftime("%d.%m.%Y") if pruvodka.drive_date else ""
    return values


def _replace_placeholders(value: str, pruvodka: Pruvodka) -> str:
    output = value
    for token, replacement in _placeholder_values(pruvodka).items():
        output = output.replace(token, replacement)
    return output


def generate_xlsx(pruvodka: Pruvodka, template_path: Path | None = None) -> Path:
    """Generate output xlsx by replacing {{...}} placeholders."""
    if load_workbook is None:
        raise RuntimeError("Knihovna openpyxl není nainstalovaná.")

    template = template_path or (settings.BASE_DIR / "pruvodka.xlsx")
    workbook = load_workbook(template)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value and "}}" in cell.value:
                    cell.value = _replace_placeholders(cell.value, pruvodka)

    target_dir = settings.MEDIA_ROOT / "pruvodky"
    target_dir.mkdir(parents=True, exist_ok=True)
    snd = (pruvodka.snd_cislo_listu or "").strip() or str(pruvodka.pk)
    safe_snd = re.sub(r"[^0-9A-Za-z._-]+", "-", snd).strip("-._")
    safe_snd = safe_snd or str(pruvodka.pk)
    output_path = target_dir / f"pruvodka_{safe_snd}.xlsx"
    if output_path.exists():
        output_path = target_dir / f"pruvodka_{safe_snd}_{uuid4().hex[:8]}.xlsx"
    workbook.save(output_path)
    return output_path


def _route_distances_osrm(source_lat: float, source_lng: float, destinations: list[tuple[float, float]]) -> list[float | None]:
    if not destinations:
        return []
    coords = [f"{source_lng:.7f},{source_lat:.7f}"] + [f"{lng:.7f},{lat:.7f}" for lat, lng in destinations]
    coords_path = ";".join(coords)
    params = urlencode(
        {
            "sources": 0,
            "annotations": "distance,duration",
        }
    )
    url = f"https://router.project-osrm.org/table/v1/driving/{coords_path}?{params}"
    req = Request(url, headers={"User-Agent": "Jedis-Django/1.0 (suffix-routing)"})
    try:
        with urlopen(req, timeout=15) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, json.JSONDecodeError):
        return [None for _ in destinations]
    distances = (payload.get("distances") or [[]])[0]
    if not isinstance(distances, list):
        return [None for _ in destinations]
    out: list[float | None] = []
    for idx in range(1, len(destinations) + 1):
        value = distances[idx] if idx < len(distances) else None
        out.append(float(value) if isinstance(value, (int, float)) else None)
    return out


def rank_suffix_sites_for_waste(
    waste_type_id: int,
    source_lat: float | None = None,
    source_lng: float | None = None,
) -> list[dict]:
    try:
        wt = WasteTypeCatalog.objects.get(pk=waste_type_id)
        waste_code = (wt.code or "").strip()
    except WasteTypeCatalog.DoesNotExist:
        wt = None
        waste_code = ""

    qs = SuffixSiteProfile.objects.select_related("suffix").exclude(suffix__isnull=True)
    if waste_code:
        qs = qs.filter(
            models.Q(accepted_waste_types__id=waste_type_id)
            | models.Q(accepted_waste_types__code=waste_code)
        )
    else:
        qs = qs.filter(accepted_waste_types__id=waste_type_id)

    sites = list(qs.distinct())
    if not sites:
        return []

    rows: list[dict] = []
    route_candidates: list[tuple[float, float]] = []
    route_indexes: list[int] = []
    for site in sites:
        row = {
            "suffix_id": site.suffix_id,
            "suffix_value": site.suffix.value,
            "site_name": site.name or "",
            "distance_m": None,
        }
        rows.append(row)
        if source_lat is not None and source_lng is not None and site.latitude is not None and site.longitude is not None:
            route_indexes.append(len(rows) - 1)
            route_candidates.append((float(site.latitude), float(site.longitude)))

    if route_candidates:
        distances = _route_distances_osrm(float(source_lat), float(source_lng), route_candidates)
        for idx, distance in zip(route_indexes, distances):
            rows[idx]["distance_m"] = distance

    rows.sort(
        key=lambda item: (
            item["distance_m"] is None,
            item["distance_m"] if item["distance_m"] is not None else 10**15,
            item["suffix_value"],
        )
    )
    return rows
