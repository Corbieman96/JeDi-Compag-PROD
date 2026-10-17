from __future__ import annotations

from io import BytesIO

from reports.services.date_range import ReportDateRange
from reports.services.suffix_workload_summary import SuffixGroupSection, SuffixWorkloadRow, SuffixWorkloadSummary

try:
    from openpyxl import Workbook
except ModuleNotFoundError:  # pragma: no cover
    Workbook = None

MAX_DETAIL_ROWS = 10_000


def _write_rows(ws, headers: list[str], rows: list[SuffixWorkloadRow], *, include_share: bool) -> None:
    ws.append(headers)
    for row in rows:
        if include_share:
            ws.append([row.label, row.count, row.share_pct])
        else:
            ws.append([row.label, row.count])


def _write_group_sheet(
    ws,
    sections: list[SuffixGroupSection],
    *,
    group_col: str,
    max_rows: int,
) -> None:
    ws.append([group_col, "Koncovka", "Počet", "Podíl %"])
    written = 0
    truncated = False
    for section in sections:
        group_label = section.title
        for row in section.suffix_rows:
            if written >= max_rows:
                truncated = True
                break
            ws.append([group_label, row.label, row.count, row.share_pct])
            written += 1
        if truncated:
            break
    if truncated:
        ws.insert_rows(1)
        ws.cell(row=1, column=1, value=f"Export omezen na {max_rows} řádků (zbytek neexportován).")


def build_suffix_workload_xlsx_bytes(summary: SuffixWorkloadSummary, date_range: ReportDateRange) -> tuple[bytes, str]:
    if Workbook is None:  # pragma: no cover
        raise RuntimeError("openpyxl není nainstalován.")

    wb = Workbook()
    ws_summary = wb.active
    assert ws_summary is not None
    ws_summary.title = "Souhrn"
    ws_summary.append(["Období od", date_range.date_from.isoformat()])
    ws_summary.append(["Období do", date_range.date_to.isoformat()])
    ws_summary.append(["Celkem jízd", summary.total])

    ws_suffix = wb.create_sheet("Koncovky")
    _write_rows(ws_suffix, ["Koncovka", "Počet", "Podíl %"], summary.by_suffix, include_share=True)

    ws_rz = wb.create_sheet("Vozidla RZ")
    _write_rows(ws_rz, ["RZ", "Počet jízd", "Podíl %"], summary.by_vehicle_rz_overview, include_share=True)

    ws_vehicle = wb.create_sheet("Koncovky u vozidla")
    _write_group_sheet(ws_vehicle, summary.by_vehicle, group_col="Vozidlo (RZ)", max_rows=MAX_DETAIL_ROWS)

    ws_waste = wb.create_sheet("Koncovky u druhu odpadu")
    _write_group_sheet(ws_waste, summary.by_waste_type, group_col="Druh odpadu", max_rows=MAX_DETAIL_ROWS)

    ws_pairs = wb.create_sheet("Koncovka a odpad")
    _write_rows(ws_pairs, ["Koncovka / odpad", "Počet", "Podíl %"], summary.by_suffix_waste_type, include_share=True)

    buf = BytesIO()
    wb.save(buf)
    filename = f"report_koncovky_{date_range.date_from.isoformat()}_{date_range.date_to.isoformat()}.xlsx"
    return buf.getvalue(), filename
