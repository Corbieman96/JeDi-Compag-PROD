from __future__ import annotations

from io import BytesIO

from reports.services.date_range import ReportDateRange
from reports.services.drivers_summary import DriversSummary, DriversSummaryRow

try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:  # pragma: no cover
    Workbook = None
    get_column_letter = None


def _write_rows(ws, headers: list[str], rows: list[DriversSummaryRow], *, include_share: bool) -> None:
    ws.append(headers)
    for row in rows:
        if include_share:
            ws.append([row.label, row.count, row.share_pct])
        else:
            ws.append([row.label, row.count])


def build_drivers_xlsx_bytes(summary: DriversSummary, date_range: ReportDateRange) -> tuple[bytes, str]:
    if Workbook is None:  # pragma: no cover
        raise RuntimeError("openpyxl není nainstalován.")

    wb = Workbook()
    ws_summary = wb.active
    assert ws_summary is not None
    ws_summary.title = "Souhrn"
    ws_summary.append(["Období od", date_range.date_from.isoformat()])
    ws_summary.append(["Období do", date_range.date_to.isoformat()])
    ws_summary.append(["Jízd s řidičem", summary.total_rides])
    ws_summary.append(["Dokončené řidičem", summary.driver_completed_count])
    ws_summary.append(["Průměr jízd / řidič-den", summary.avg_rides_per_driver_day])
    ws_summary.append(["Řidič s jízdou i absencí", summary.drivers_with_absence_and_rides])
    ws_summary.append(["Popelář — dny", float(summary.popelar.days)])
    ws_summary.append(["Popelář — hodiny", float(summary.popelar.hours)])
    ws_summary.append(["Popelář — sazba Kč/h", float(summary.popelar.rate_czk)])
    ws_summary.append(["Popelář — celkem Kč", float(summary.popelar.total_czk)])

    ws_rides = wb.create_sheet("Jízdy na řidiče")
    _write_rows(ws_rides, ["Řidič", "Počet", "Podíl %"], summary.rides_by_driver, include_share=True)

    ws_abs_type = wb.create_sheet("Nepřítomnosti podle typu")
    _write_rows(ws_abs_type, ["Typ", "Počet", "Podíl %"], summary.absences_by_type, include_share=True)

    ws_abs_driver = wb.create_sheet("Nepřítomnosti podle řidiče")
    _write_rows(ws_abs_driver, ["Řidič / typ", "Počet", "Podíl %"], summary.absences_by_driver, include_share=True)

    if summary.absences_other:
        ws_other = wb.create_sheet("Jiné")
        ws_other.append(["Řidič", "Období", "Dní", "Poznámka"])
        for row in summary.absences_other:
            ws_other.append([row.driver_label, row.range_label, row.days, row.note or ""])

    ws_day_rz = wb.create_sheet("Den RZ")
    day_rz_headers = ["Řidič", "Den", "Druh vozidla", "RZ"]
    ws_day_rz.append(day_rz_headers)
    for row in summary.day_rz_rows:
        ws_day_rz.append([row.driver_label, row.date_label, row.vehicle_type, row.rz])
    last_row = max(1, ws_day_rz.max_row)
    ws_day_rz.auto_filter.ref = f"A1:D{last_row}"
    col_widths = [len(h) for h in day_rz_headers]
    for row in summary.day_rz_rows:
        values = (row.driver_label, row.date_label, row.vehicle_type, row.rz)
        for i, val in enumerate(values):
            col_widths[i] = max(col_widths[i], len(str(val or "")))
    for i, width in enumerate(col_widths, start=1):
        # openpyxl šířka ≈ znaky + drobná rezerva na filtr/okraje
        ws_day_rz.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 48)

    ws_popelar = wb.create_sheet("Popelář")
    ws_popelar.append(["Dny", float(summary.popelar.days)])
    ws_popelar.append(["Hodiny za den", float(summary.popelar.hours_per_day)])
    ws_popelar.append(["Hodiny celkem", float(summary.popelar.hours)])
    ws_popelar.append(["Sazba Kč/h", float(summary.popelar.rate_czk)])
    ws_popelar.append(["Celkem Kč", float(summary.popelar.total_czk)])

    buf = BytesIO()
    wb.save(buf)
    filename = f"report_ridici_{date_range.date_from.isoformat()}_{date_range.date_to.isoformat()}.xlsx"
    return buf.getvalue(), filename
