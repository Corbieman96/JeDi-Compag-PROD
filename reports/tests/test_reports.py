from __future__ import annotations

import uuid
from datetime import date

from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from django.urls import reverse

from catalogs.models import (
    AddressCatalog,
    DriverAbsence,
    DriverProfile,
    IcoCatalog,
    NameCatalog,
    WasteCategoryCatalog,
    WasteCodeCatalog,
    WasteTypeCatalog,
)
from core.models import AppSettings
from reports.services.date_range import ReportDateRange
from reports.services.customer_requests_summary import build_customer_requests_summary
from reports.services.drivers_summary import build_drivers_summary
from reports.services.exports.drivers_export import build_drivers_xlsx_bytes
from reports.services.exports.suffix_workload_export import build_suffix_workload_xlsx_bytes
from reports.services.hub_summary import build_hub_summary
from reports.services.operations_summary import build_operations_summary
from reports.services.print_summary import build_print_summary
from reports.services.rides_summary import build_rides_summary
from reports.services.suffix_workload_summary import build_suffix_workload_summary
from reports.services.svoz_summary import build_svoz_summary
from users.constants import ROLE_DISPECER
from users.models import UserAccessProfile
from users.services.access_profile import get_access_profile
from users.services.nav_visibility import effective_show_reports_nav
from users.services.roles import set_user_role
from waste_docs.models import Pruvodka

User = get_user_model()


def _minimal_pruvodka(*, created_by, drive_date: date, status=Pruvodka.Status.RECORDED, driver=None):
    suffix = uuid.uuid4().hex[:8]
    name = NameCatalog.objects.create(value=f"Subjekt R {suffix}")
    ico = IcoCatalog.objects.create(value=f"{suffix[:8]}".ljust(8, "0")[:8])
    addr = AddressCatalog.objects.create(value=f"Testovací ulice {suffix}")
    wt = WasteTypeCatalog.objects.create(code=f"WT{suffix[:6]}", value=f"Typ R {suffix}")
    wc = WasteCodeCatalog.objects.create(value=f"C{suffix[:4]}")
    wcat = WasteCategoryCatalog.objects.create(value=f"K{suffix[:4]}")
    return Pruvodka.objects.create(
        name=name,
        ico=ico,
        address=addr,
        waste_type=wt,
        waste_code=wc,
        waste_category=wcat,
        created_by=created_by,
        drive_date=drive_date,
        status=status,
        driver=driver,
    )


class ReportsNavVisibilityTests(TestCase):
    def test_superuser_always_sees_reports_nav(self):
        u = User.objects.create_user(username="rep_su", password="x", is_superuser=True)
        self.assertTrue(effective_show_reports_nav(u))

    def test_dispecer_inherit_default_false(self):
        u = User.objects.create_user(username="rep_disp", password="x")
        set_user_role(u, ROLE_DISPECER)
        prof = get_access_profile(u)
        prof.nav_reports = UserAccessProfile.NavOverride.INHERIT
        prof.save(update_fields=["nav_reports"])
        self.assertFalse(effective_show_reports_nav(u))

    def test_dispecer_show_override(self):
        u = User.objects.create_user(username="rep_disp_show", password="x")
        set_user_role(u, ROLE_DISPECER)
        prof = get_access_profile(u)
        prof.nav_reports = UserAccessProfile.NavOverride.SHOW
        prof.save(update_fields=["nav_reports"])
        self.assertTrue(effective_show_reports_nav(u))

    def test_role_matrix_can_enable_reports(self):
        u = User.objects.create_user(username="rep_matrix", password="x")
        set_user_role(u, ROLE_DISPECER)
        prof = get_access_profile(u)
        prof.nav_reports = UserAccessProfile.NavOverride.INHERIT
        prof.save(update_fields=["nav_reports"])
        s = AppSettings.load()
        s.role_nav_defaults = {ROLE_DISPECER: {"reports": True}}
        s.save(update_fields=["role_nav_defaults"])
        self.assertTrue(effective_show_reports_nav(u))


class ReportsViewSmokeTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username="rep_admin", password="pw", is_superuser=True)
        self.dispecer = User.objects.create_user(username="rep_disp2", password="pw")
        set_user_role(self.dispecer, ROLE_DISPECER)

    def test_index_200_for_superuser(self):
        c = Client()
        c.force_login(self.admin)
        r = c.get(reverse("reports:index"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Reporty")
        self.assertContains(r, "Jízdy v období")

    def test_operations_report_200(self):
        c = Client()
        c.force_login(self.admin)
        r = c.get(reverse("reports:operations"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Provoz")

    def test_index_redirects_unauthorized(self):
        c = Client()
        c.force_login(self.dispecer)
        r = c.get(reverse("reports:index"))
        self.assertEqual(r.status_code, 302)


class RidesSummaryTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="rep_rides", password="x")
        self.driver = DriverProfile.objects.create(personal_number="RR-1", last_name="Novák", first_name="Jan")
        self.d1 = date(2026, 6, 1)
        self.d2 = date(2026, 6, 2)
        _minimal_pruvodka(created_by=self.user, drive_date=self.d1, status=Pruvodka.Status.RECORDED)
        _minimal_pruvodka(created_by=self.user, drive_date=self.d1, status=Pruvodka.Status.GENERATED, driver=self.driver)
        _minimal_pruvodka(created_by=self.user, drive_date=self.d2, status=Pruvodka.Status.RECORDED)

    def test_build_rides_summary_counts(self):
        summary = build_rides_summary(ReportDateRange(date_from=self.d1, date_to=self.d2))
        self.assertEqual(summary.total, 3)
        self.assertEqual(len(summary.by_status), 2)
        self.assertEqual(len(summary.by_day), 2)
        status_map = {row.label: row.count for row in summary.by_status}
        self.assertEqual(status_map.get("Založeno"), 2)
        self.assertEqual(status_map.get("Vygenerováno"), 1)
        self.assertGreaterEqual(len(summary.by_activity), 0)

    def test_rides_extended_kpis(self):
        ride = Pruvodka.objects.filter(drive_date=self.d1).first()
        ride.activity_snapshot = "Svoz"
        ride.follow_up_open = True
        ride.save(update_fields=["activity_snapshot", "follow_up_open"])
        summary = build_rides_summary(ReportDateRange(date_from=self.d1, date_to=self.d2))
        self.assertEqual(summary.follow_up_open_count, 1)
        activity_map = {row.label: row.count for row in summary.by_activity}
        self.assertEqual(activity_map.get("Svoz"), 1)


class SuffixWorkloadSummaryTests(TestCase):
    def setUp(self):
        from catalogs.models import SuffixCatalog, VehicleRzCatalog

        self.user = User.objects.create_user(username="rep_suffix", password="x")
        self.d1 = date(2026, 6, 10)
        suffix_a = SuffixCatalog.objects.create(value="Spalovna")
        suffix_b = SuffixCatalog.objects.create(value="Skládka")
        rz_a = VehicleRzCatalog.objects.create(value="1A11111")
        rz_b = VehicleRzCatalog.objects.create(value="2B22222")
        ride_a = _minimal_pruvodka(created_by=self.user, drive_date=self.d1)
        ride_a.suffix = suffix_a
        ride_a.vehicle_rz = rz_a
        ride_a.waste_type_snapshot = "Komunální"
        ride_a.activity_snapshot = "Svoz"
        ride_a.save()
        ride_b = _minimal_pruvodka(created_by=self.user, drive_date=self.d1)
        ride_b.suffix = suffix_b
        ride_b.vehicle_rz = rz_b
        ride_b.waste_type_snapshot = "Stavební"
        ride_b.activity_snapshot = "Svoz"
        ride_b.save()

    def test_build_suffix_workload_summary(self):
        summary = build_suffix_workload_summary(ReportDateRange(date_from=self.d1, date_to=self.d1))
        self.assertEqual(summary.total, 2)
        suffix_map = {row.label: row.count for row in summary.by_suffix}
        self.assertEqual(suffix_map.get("Spalovna"), 1)
        self.assertEqual(suffix_map.get("Skládka"), 1)
        self.assertEqual(len(summary.by_vehicle), 2)
        vehicle_titles = {group.title for group in summary.by_vehicle}
        self.assertIn("Vozidlo 1A11111", vehicle_titles)
        self.assertIn("Vozidlo 2B22222", vehicle_titles)
        group_a = next(g for g in summary.by_vehicle if "1A11111" in g.title)
        self.assertEqual(group_a.total_rides, 1)
        self.assertEqual(len(group_a.suffix_rows), 1)
        self.assertEqual(group_a.suffix_rows[0].label, "Spalovna")

    def test_vehicle_rz_overview(self):
        summary = build_suffix_workload_summary(ReportDateRange(date_from=self.d1, date_to=self.d1))
        rz_map = {row.label: row.count for row in summary.by_vehicle_rz_overview}
        self.assertEqual(rz_map.get("1A11111"), 1)
        self.assertEqual(rz_map.get("2B22222"), 1)
        self.assertEqual(summary.by_vehicle_rz_overview[0].share_pct, 50.0)

    def test_suffix_workload_xlsx_export(self):
        summary = build_suffix_workload_summary(ReportDateRange(date_from=self.d1, date_to=self.d1))
        payload, filename = build_suffix_workload_xlsx_bytes(summary, ReportDateRange(date_from=self.d1, date_to=self.d1))
        self.assertTrue(payload.startswith(b"PK"))
        self.assertIn("report_koncovky", filename)
        from openpyxl import load_workbook
        from io import BytesIO

        wb = load_workbook(BytesIO(payload))
        self.assertIn("Souhrn", wb.sheetnames)
        self.assertIn("Vozidla RZ", wb.sheetnames)
        self.assertIn("Koncovky u vozidla", wb.sheetnames)
        self.assertIn("Koncovka a odpad", wb.sheetnames)

    def test_suffix_waste_type_pairs(self):
        summary = build_suffix_workload_summary(ReportDateRange(date_from=self.d1, date_to=self.d1))
        self.assertEqual(len(summary.by_suffix_waste_type), 2)


class DriversExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="rep_drivers", password="x")
        self.driver = DriverProfile.objects.create(personal_number="DR-1", last_name="Svoboda", first_name="Petr")
        self.d1 = date(2026, 6, 15)
        ride = _minimal_pruvodka(
            created_by=self.user,
            drive_date=self.d1,
            status=Pruvodka.Status.GENERATED,
            driver=self.driver,
        )
        ride.driver_completed_at = ride.created_at
        ride.save(update_fields=["driver_completed_at"])
        DriverAbsence.objects.create(
            driver=self.driver,
            absence_type=DriverAbsence.AbsenceType.VACATION,
            date_from=self.d1,
            date_to=self.d1,
        )

    def test_drivers_xlsx_export(self):
        date_range = ReportDateRange(date_from=self.d1, date_to=self.d1)
        summary = build_drivers_summary(date_range)
        payload, filename = build_drivers_xlsx_bytes(summary, date_range)
        self.assertTrue(payload.startswith(b"PK"))
        self.assertIn("report_ridici", filename)
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(payload))
        self.assertIn("Souhrn", wb.sheetnames)
        self.assertIn("Jízdy na řidiče", wb.sheetnames)
        self.assertIn("Nepřítomnosti podle typu", wb.sheetnames)
        self.assertIn("Nepřítomnosti podle řidiče", wb.sheetnames)
        self.assertIn("Popelář", wb.sheetnames)
        self.assertIn("Den RZ", wb.sheetnames)
        self.assertEqual(wb["Souhrn"]["A3"].value, "Jízd s řidičem")
        self.assertEqual(wb["Souhrn"]["B3"].value, 1)


class DriversDayRzTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="rep_day_rz", password="x")
        self.driver = DriverProfile.objects.create(
            personal_number="DR-RZ",
            last_name="Poliščuk",
            first_name="Jan",
        )
        self.d1 = date(2026, 9, 5)

    def test_two_rz_same_day_two_rows(self):
        from catalogs.models import VehicleKindCatalog, VehicleRzCatalog, VehicleRzProfile

        kind_retez = VehicleKindCatalog.objects.create(name="Řetěz", shortcut="R-dayrz", sort_order=1)
        kind_hak = VehicleKindCatalog.objects.create(name="Hák", shortcut="H-dayrz", sort_order=2)

        rz_a = VehicleRzCatalog.objects.create(value="1SD1235")
        rz_b = VehicleRzCatalog.objects.create(value="2SD9999")
        prof_a = VehicleRzProfile.objects.create(rz=rz_a, vehicle_type="poznámka nakladač")
        prof_a.vehicle_kinds.add(kind_retez)
        prof_b = VehicleRzProfile.objects.create(rz=rz_b, vehicle_type="jiná poznámka")
        prof_b.vehicle_kinds.add(kind_hak)

        ride_a = _minimal_pruvodka(
            created_by=self.user,
            drive_date=self.d1,
            status=Pruvodka.Status.GENERATED,
            driver=self.driver,
        )
        ride_a.vehicle_rz = rz_a
        ride_a.save(update_fields=["vehicle_rz"])

        ride_b = _minimal_pruvodka(
            created_by=self.user,
            drive_date=self.d1,
            status=Pruvodka.Status.GENERATED,
            driver=self.driver,
        )
        ride_b.vehicle_rz = rz_b
        ride_b.save(update_fields=["vehicle_rz"])

        # Same RZ twice same day → still one row for that RZ.
        ride_c = _minimal_pruvodka(
            created_by=self.user,
            drive_date=self.d1,
            status=Pruvodka.Status.GENERATED,
            driver=self.driver,
        )
        ride_c.vehicle_rz = rz_a
        ride_c.save(update_fields=["vehicle_rz"])

        summary = build_drivers_summary(ReportDateRange(date_from=self.d1, date_to=self.d1))
        self.assertEqual(len(summary.day_rz_rows), 2)
        by_rz = {r.rz: r for r in summary.day_rz_rows}
        self.assertIn("1SD1235", by_rz)
        self.assertIn("2SD9999", by_rz)
        self.assertIn("Poliščuk", by_rz["1SD1235"].driver_label)
        self.assertEqual(by_rz["1SD1235"].date_label, "5. 9. 2026")
        self.assertEqual(by_rz["1SD1235"].vehicle_type, "Řetěz")
        self.assertEqual(by_rz["2SD9999"].vehicle_type, "Hák")

        payload, _filename = build_drivers_xlsx_bytes(
            summary, ReportDateRange(date_from=self.d1, date_to=self.d1)
        )
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(payload))
        self.assertIn("Den RZ", wb.sheetnames)
        self.assertEqual(wb["Den RZ"]["A1"].value, "Řidič")
        self.assertEqual(wb["Den RZ"]["C1"].value, "Druh vozidla")
        self.assertEqual(wb["Den RZ"].max_row, 3)

class DriversOtherAbsenceTests(TestCase):
    def setUp(self):
        self.driver = DriverProfile.objects.create(
            personal_number="DR-OTH",
            last_name="Horák",
            first_name="Karel",
        )
        self.d1 = date(2026, 8, 10)
        self.d2 = date(2026, 8, 12)

    def test_absences_other_includes_note(self):
        DriverAbsence.objects.create(
            driver=self.driver,
            absence_type=DriverAbsence.AbsenceType.OTHER,
            date_from=self.d1,
            date_to=self.d2,
            note="Školení mimo katalog",
        )
        DriverAbsence.objects.create(
            driver=self.driver,
            absence_type=DriverAbsence.AbsenceType.VACATION,
            date_from=self.d1,
            date_to=self.d1,
            note="nemá být v Jiné",
        )
        summary = build_drivers_summary(ReportDateRange(date_from=self.d1, date_to=self.d2))
        self.assertEqual(len(summary.absences_other), 1)
        row = summary.absences_other[0]
        self.assertIn("Horák", row.driver_label)
        self.assertEqual(row.note, "Školení mimo katalog")
        self.assertEqual(row.days, 3)

    def test_absences_other_empty_without_other_type(self):
        DriverAbsence.objects.create(
            driver=self.driver,
            absence_type=DriverAbsence.AbsenceType.SICK,
            date_from=self.d1,
            date_to=self.d1,
        )
        summary = build_drivers_summary(ReportDateRange(date_from=self.d1, date_to=self.d1))
        self.assertEqual(summary.absences_other, [])

    def test_drivers_xlsx_includes_jine_sheet(self):
        DriverAbsence.objects.create(
            driver=self.driver,
            absence_type=DriverAbsence.AbsenceType.OTHER,
            date_from=self.d1,
            date_to=self.d1,
            note="Důvod X",
        )
        date_range = ReportDateRange(date_from=self.d1, date_to=self.d1)
        summary = build_drivers_summary(date_range)
        payload, _filename = build_drivers_xlsx_bytes(summary, date_range)
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(payload))
        self.assertIn("Jiné", wb.sheetnames)
        self.assertEqual(wb["Jiné"]["D2"].value, "Důvod X")


class PopelarCostTests(TestCase):
    def setUp(self):
        self.driver = DriverProfile.objects.create(
            personal_number="DR-POP",
            last_name="Novák",
            first_name="Jan",
        )
        self.d1 = date(2026, 7, 1)
        self.d2 = date(2026, 7, 2)
        settings = AppSettings.load()
        settings.popelar_hours_per_day = "8.00"
        settings.popelar_hourly_rate_czk = "250.00"
        settings.save(update_fields=["popelar_hours_per_day", "popelar_hourly_rate_czk"])

    def test_full_and_partial_popelar_days(self):
        from decimal import Decimal

        DriverAbsence.objects.create(
            driver=self.driver,
            absence_type=DriverAbsence.AbsenceType.GARBAGE,
            date_from=self.d1,
            date_to=self.d1,
            is_partial=False,
        )
        DriverAbsence.objects.create(
            driver=self.driver,
            absence_type=DriverAbsence.AbsenceType.GARBAGE,
            date_from=self.d2,
            date_to=self.d2,
            is_partial=True,
        )
        summary = build_drivers_summary(ReportDateRange(date_from=self.d1, date_to=self.d2))
        self.assertEqual(summary.popelar.days, Decimal("1.50"))
        self.assertEqual(summary.popelar.hours, Decimal("12.00"))
        self.assertEqual(summary.popelar.rate_czk, Decimal("250.00"))
        self.assertEqual(summary.popelar.total_czk, Decimal("3000.00"))


class HubSummaryTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="rep_hub", password="x")
        self.d1 = date(2026, 6, 20)

    def test_build_hub_summary(self):
        _minimal_pruvodka(created_by=self.user, drive_date=self.d1, status=Pruvodka.Status.GENERATED)
        hub = build_hub_summary(ReportDateRange(date_from=self.d1, date_to=self.d1))
        self.assertEqual(hub.rides_total, 1)
        self.assertEqual(hub.rides_generated_pct, 100.0)


class OperationsSummaryTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="rep_ops", password="x")
        self.d1 = date(2026, 6, 21)
        ride = _minimal_pruvodka(created_by=self.user, drive_date=self.d1)
        ride.activity_snapshot = "Svoz"
        ride.save(update_fields=["activity_snapshot"])

    def test_build_operations_summary(self):
        summary = build_operations_summary(ReportDateRange(date_from=self.d1, date_to=self.d1))
        self.assertGreaterEqual(len(summary.kpis), 10)
        self.assertGreaterEqual(len(summary.rides_by_activity), 1)


class PrintSummaryTests(TestCase):
    def setUp(self):
        from django.utils import timezone
        from waste_docs.models import PrintFileMark

        self.user = User.objects.create_user(username="rep_print", password="x")
        self.d1 = timezone.localdate()
        PrintFileMark.objects.create(user=self.user, drive_date=self.d1, row_key="test-row")

    def test_print_share_pct(self):
        summary = build_print_summary(ReportDateRange(date_from=self.d1, date_to=self.d1))
        self.assertEqual(summary.total_marks, 1)
        self.assertEqual(summary.by_user[0].share_pct, 100.0)
