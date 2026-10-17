from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from catalogs.models import WasteTypeCatalog


@dataclass
class WasteTypeImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0


def _normalize_code(raw: str) -> str:
    return " ".join((raw or "").strip().split())


def import_waste_types_from_xlsx(path: Path) -> WasteTypeImportResult:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    result = WasteTypeImportResult()

    for row in ws.iter_rows(values_only=True):
        code = _normalize_code(str(row[1] or ""))
        name = str(row[2] or "").strip()
        category = str(row[3] or "").strip()

        if not code or not name:
            result.skipped += 1
            continue

        obj, created = WasteTypeCatalog.objects.get_or_create(
            value=name,
            defaults={"code": code, "category": category},
        )
        if created:
            result.created += 1
            continue

        changed = False
        if obj.code != code:
            obj.code = code
            changed = True
        if obj.category != category:
            obj.category = category
            changed = True
        if changed:
            obj.save(update_fields=("code", "category"))
            result.updated += 1
        else:
            result.skipped += 1

    return result
