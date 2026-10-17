"""Jednorázové vytvoření pio/static/pio/pio_template.xlsx (spustit z kořene projektu)."""

from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "pio" / "static" / "pio" / "pio_template.xlsx"


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "PIO"
    ws["A1"] = "Prvotní informace o odpadu — předvyplnění (JeDiS)"
    ws["A2"] = "Sloupec B obsahuje tokeny {{PIO_*}}; šablonu můžete nahradit vlastní (viz settings PIO_XLSX_TEMPLATE)."
    rows = [
        ("Obchodní firma / název", "{{PIO_COMPANY}}"),
        ("Název F/P osoby", "{{PIO_NAME_FP}}"),
        ("IČO", "{{PIO_ICO}}"),
        ("Adresa plnění", "{{PIO_FULFILLMENT_ADDRESS}}"),
        ("Adresa provozovny", "{{PIO_FACILITY_ADDRESS}}"),
        ("ORP (plnění)", "{{PIO_ORP}}"),
        ("IČZÚJ (plnění)", "{{PIO_ICZUJ}}"),
        ("Kód odpadu", "{{PIO_WASTE_CODE}}"),
        ("Název druhu odpadu", "{{PIO_WASTE_NAME}}"),
        ("Kategorie odpadu", "{{PIO_WASTE_CATEGORY}}"),
        ("Koncovka", "{{PIO_SUFFIX}}"),
        ("Číslo smlouvy (pár)", "{{PIO_CONTRACT}}"),
        ("Datum", "{{PIO_DATE}}"),
    ]
    start = 4
    for i, (label, token) in enumerate(rows):
        ws.cell(row=start + i, column=1, value=label)
        ws.cell(row=start + i, column=2, value=token)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 48
    wb.save(OUT)
    print("Uloženo:", OUT)


if __name__ == "__main__":
    main()
